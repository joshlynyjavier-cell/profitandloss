import openpyxl, json, sys
import matplotlib; matplotlib.use('svg')
import matplotlib.pyplot as plt
from matplotlib import font_manager as fm
for f in fm.findSystemFonts(): 
    if 'Inter' in f: fm.fontManager.addfont(f)
plt.rcParams.update({'font.family':'Inter','font.size':10,'axes.spines.top':False,'axes.spines.right':False,
  'axes.edgecolor':'#9AA3AE','axes.labelcolor':'#3B4452','xtick.color':'#59636F','ytick.color':'#59636F',
  'axes.grid':True,'grid.color':'#E6E9ED','grid.linewidth':0.8,'axes.axisbelow':True,'svg.fonttype':'none'})
NAVY='#1F3A5F'; GRAY='#9AA3AE'; LG='#C9D0D8'; RED='#B5482F'; GRN='#2E7D5B'; INK='#1E2530'
wb=openpyxl.load_workbook(sys.argv[1],data_only=True)
g=lambda sh,c: wb[sh][c].value
V={}
# ---------- values used in deck ----------
pl=wb['PL_Waterfall']; pm=wb['Platform_Market']; es=wb['Exec_Summary']; dd=wb['DoorDash']; pe=wb['Promo_Escalation']; inc=wb['Incrementality']; mk=wb['Markup']
for k,c in dict(gross_fy='E6',net_fy='E13',cm1_fy='E16',cm2_fy='E18',gross_exit='F6',cm2_exit='F18',orders_fy='E5').items(): V[k]=pl[c].value
for k,c in dict(take_fy='C14',take_exit='G14',cm2p_fy='C18',cm2p_exit='G18',be_fy='C23',be_exit='G23',share_dinein='C22',cm2_order='D18',aov='D6').items(): V[k]=pl[c].value
V['wf']={pl[f'A{r}'].value:pl[f'E{r}'].value for r in range(6,19)}
# segments
hdrs=[pm.cell(row=4,column=c).value for c in range(2,20)]
seg={}
for i,h in enumerate(hdrs):
    col=i+2
    seg[h]={'cm2p':pm.cell(row=26,column=col).value,'exit_cm2p':pm.cell(row=37,column=col).value,'gross_ann':pm.cell(row=33,column=col).value,
            'cm2_ann':pm.cell(row=34,column=col).value,'take':pm.cell(row=18,column=col).value,'promo_p':pm.cell(row=13,column=col).value/pm.cell(row=10,column=col).value,
            'cm2_order':pm.cell(row=28,column=col).value,'rev_mix':pm.cell(row=29,column=col).value,'cm_mix':pm.cell(row=30,column=col).value,'be':pm.cell(row=32,column=col).value,
            'exit_promo':pm.cell(row=36,column=col).value}
V['seg']=seg
V['pe']={k:pe[c].value for k,c in dict(pre_promo='B5',exit_promo='C5',dd_pre='B6',dd_exit='C6',ue_exit='C7',ec_exit='C8',gpsw_pre='B9',gpsw_exit='C9',gpsw_chg='D9',opsw_chg='D10',cm2_pre='B12',cm2_exit='C12',ratio_pre='B16',ratio_post='B17',esc_cost='B22',rollback='B25',ratio_dd='B27').items()}
V['inc']={k:inc[c].value for k,c in dict(g_in='D21',g_pick='E21',g_3pd='G21',g_ctrl='B24',disp='B29',incr='B30',corr='B32',peak='H43',peak_util='I43',off_util='J43',r90_dine='E49',r90_1p='E50',r90_dd='E51',r90_ddp='E52',r90_ue='E53',r90_uep='E54',r90_gh='E55',share_q1='J19',share_q4='J20').items()}
V['inc']['ev']={lab:[inc[f'{c}{r}'].value for c in 'BCDEF'] for lab,r in [('bear',78),('base',79),('bull',80),('base_exit',83)]}
V['inc']['gap']={inc[f'A{r}'].value:inc[f'D{r}'].value for r in range(59,65)}
V['mk']={k:mk[c].value for k,c in dict(v5='C10',v10='D10',v15='E10',n5='C12',n10='D12',n15='E12',t5='C13',el='B16',r2='B17',implied='E31',banked='E32',mech='E34').items()}
V['dd']={k:dd[c].value for k,c in dict(g1='B6',g2='B7',ad_share='B11',cm2p='B12',acc12='C25',acc24='C26',top12='D25',top24='D26',ctr12='E25',ctr24='E26',walk12='F25',walk24='F26',be_sn='B33',be_tu='B34',sav_sn='B35',sav_tu='B36',mk_off='B37',post_topup='B39',cm2_acc='C27',cm2_top='D27',cm2_ctr='E27',take_acc='C28',take_top='D28',mand='D21').items()}
V['es']={k:es[c].value for k,c in dict(p1='B35',p2='B36',p3='B37',tot='B38',cm2_after='B39',cm2p_after='B40',gv_tot='B46',gv_gap='B47').items()}
# weekly
W=wb['Weeks']; weeks=[W[f'A{r}'].value for r in range(6,57)]
promo=[W[f'R{r}'].value for r in range(6,57)]; egsw=[W[f'X{r}'].value for r in range(6,57)]; cm2p=[W[f'S{r}'].value for r in range(6,57)]
# monthly channel
months=[inc[f'A{r}'].value for r in range(6,18)]
ch={k:[inc[f'{c}{r}'].value for r in range(6,18)] for k,c in dict(instore='D',pickup='E',p1='F',tpd='G',cater='H').items()}
json.dump(V,open('deck/values.json','w'),indent=1,default=str)

def save(fig,name): fig.savefig(f'deck/{name}.svg',bbox_inches='tight',transparent=True); plt.close(fig)
M=lambda x:x/1e6
# ---- C1 bridge
steps=[('FY run-rate',M(V['cm2_fy']),'tot'),('Exit run-rate\n(Jun period)',M(V['cm2_exit']),'tot'),('Roll back\npromo',M(V['es']['p1']),'d'),('5% menu\nmarkup',M(V['es']['p2']),'d'),('DoorDash\ncounter',M(V['es']['p3']),'d'),('Target\nrun-rate',M(V['es']['cm2_after']),'tot')]
fig,ax=plt.subplots(figsize=(6.4,3.3)); run=0
for i,(l,v,t) in enumerate(steps):
    if t=='tot': ax.bar(i,v,color=NAVY if i!=1 else RED,width=0.62); run=v; y=v
    else:
        b=run if v>0 else run+v; ax.bar(i,abs(v),bottom=b,color=GRN if v>0 else GRAY,width=0.62); run+=v; y=run if v>0 else b+abs(v)
    ax.text(i,y+0.25,f"{'+' if t=='d' and v>0 else ''}{'−' if v<0 else ''}${abs(v):.1f}M",ha='center',va='bottom',fontsize=9.5,color=INK,fontweight='bold')
ax.set_xticks(range(len(steps))); ax.set_xticklabels([s[0] for s in steps],fontsize=8.5); ax.set_ylim(0,19.5)
ax.set_ylabel('Annual CM2 ($M)'); ax.grid(axis='x',visible=False); save(fig,'c1_bridge')
# ---- C2 waterfall
wf=V['wf']; items=[('Gross',wf['Gross subtotal'],'tot'),('Commis-\nsion',wf['Less: platform commission'],'d'),('Mktg\ncomm.',wf['Less: marketing commission'],'d'),
 ('Promo',wf['Less: restaurant-funded promotions'],'d'),('Ads',wf['Less: advertising'],'d'),('Fees +\nrefunds',wf['Less: service fees']+wf['Less: refunds & adjustments'],'d'),
 ('Net\ndeposit',wf['Net revenue / bank deposit'],'tot'),('COGS\n30%',wf['Less: food & packaging COGS'],'d'),('CM1',wf['CM1'],'tot'),('Labor\n16%',wf['Less: variable delivery labor'],'d'),('CM2',wf['CM2 / contribution margin'],'tot')]
fig,ax=plt.subplots(figsize=(7.2,3.4)); run=0
for i,(l,v,t) in enumerate(items):
    v=M(v)
    if t=='tot': ax.bar(i,v,color=NAVY if l!='CM2' else GRN,width=0.64); run=v; ax.text(i,v+1,f"${v:.1f}M",ha='center',fontsize=8.5,color=INK,fontweight='bold')
    else: ax.bar(i,-v,bottom=run+v,color=LG if l not in('Promo',) else RED,width=0.64); ax.text(i,run+0.8,f"−{-v:.1f}",ha='center',fontsize=8,color='#59636F'); run+=v
ax.set_xticks(range(len(items))); ax.set_xticklabels([x[0] for x in items],fontsize=8); ax.set_ylabel('FY annualized ($M)'); ax.grid(axis='x',visible=False); ax.set_ylim(0,82)
save(fig,'c2_waterfall')
# ---- C3a channel index
fig,ax=plt.subplots(figsize=(5.0,2.9)); x=range(12)
own=[ch['pickup'][i]+ch['p1'][i]+ch['cater'][i] for i in x]
for s,lab,col,lw in [(ch['tpd'],'3PD gross',RED,2.2),(own,'Pickup + 1P + catering',GRAY,1.8),(ch['instore'],'In-store (dine-in)',NAVY,2.2)]:
    idx=[v/s[0]*100 for v in s]; ax.plot(x,idx,color=col,lw=lw); ax.text(11.2,idx[-1],f"{lab}  {idx[-1]:.0f}",va='center',fontsize=8.5,color=INK)
ax.set_xticks(list(x)[::2]); ax.set_xticklabels(['Jul','Sep','Nov','Jan','Mar','May'],fontsize=8.5); ax.set_xlim(-0.3,15.5)
ax.set_ylabel('Per store-week, Jul-25 = 100'); ax.spines['bottom'].set_visible(True); save(fig,'c3_trend')
# ---- C3b breakeven incrementality
labs=['Grubhub','DoorDash','All 3PD','UberEats','East Coast']; keys=['Grubhub','DoorDash','Total','UberEats','East Coast (New)']
fig,ax=plt.subplots(figsize=(4.4,2.9))
fy=[seg[k]['be'] for k in keys]; ex=[1-seg[k]['exit_cm2p']/0.48 for k in keys]
y=range(len(labs))
ax.barh([i+0.19 for i in y],[v*100 for v in fy],height=0.36,color=NAVY,label='FY')
ax.barh([i-0.19 for i in y],[v*100 for v in ex],height=0.36,color=LG,label='Exit (Jun)')
lo,hi=V['inc']['incr']*100,55
ax.axvspan(lo-0.5,hi+0.5,color=RED,alpha=0.15,lw=0); ax.text((lo+hi)/2,4.75,f'Evidence\n≈{lo:.0f}–{hi:.0f}%',ha='center',fontsize=8,color=RED,va='bottom')
for i,(a,b) in enumerate(zip(fy,ex)): ax.text(a*100+1,i+0.19,f"{a*100:.0f}%",va='center',fontsize=8); ax.text(b*100+1,i-0.19,f"{b*100:.0f}%",va='center',fontsize=8,color='#59636F')
ax.set_yticks(list(y)); ax.set_yticklabels(labs,fontsize=9); ax.set_xlim(0,100); ax.set_ylim(-0.6,5.6); ax.set_xlabel('Break-even incrementality (%)')
ax.legend(frameon=False,fontsize=8,loc='upper center',bbox_to_anchor=(0.5,-0.2),ncol=2); ax.grid(axis='y',visible=False); save(fig,'c3_breakeven')
# ---- C4 DoorDash
opts=['Counter\n28.5% / 12-mo','Accept 31%\nspend-neutral','Accept 31%\n+ ads top-up','Walk\n(base recapture)']
v12=[V['dd'][k] for k in ['ctr12','acc12','top12','walk12']]; v24=[V['dd'][k] for k in ['ctr24','acc24','top24','walk24']]
fig,ax=plt.subplots(figsize=(5.6,3.0)); x=range(4)
ax.bar([i-0.19 for i in x],[M(v) for v in v12],0.36,color=LG,label='12 months'); ax.bar([i+0.19 for i in x],[M(v) for v in v24],0.36,color=[NAVY,NAVY,NAVY,GRAY],label='24 months')
for i,(a,b) in enumerate(zip(v12,v24)): ax.text(i-0.19,M(a)+0.1,f"${M(a):.1f}M",ha='center',fontsize=8,color='#59636F'); ax.text(i+0.19,M(b)+0.1,f"${M(b):.1f}M",ha='center',fontsize=8.5,fontweight='bold',color=INK)
ax.set_xticks(list(x)); ax.set_xticklabels(opts,fontsize=8.5); ax.set_ylabel('CM2 cost vs today ($M)'); ax.legend(frameon=False,fontsize=8,loc='upper left'); ax.grid(axis='x',visible=False); ax.set_ylim(0,7.8)
save(fig,'c4_dd')
# ---- C5 promo vs volume (two stacked panels, shared x)
fig,(a1,a2)=plt.subplots(2,1,figsize=(6.0,3.5),sharex=True,gridspec_kw={'height_ratios':[1,1],'hspace':0.18})
x=range(len(weeks)); a1.plot(x,[p*100 for p in promo],color=RED,lw=2.2); a1.set_ylabel('Promo % of gross',fontsize=8.5); a1.set_ylim(0,8.5)
a1.axvline(30.5,color=GRAY,lw=1,ls='--'); a1.text(30,6.6,'Escalation begins (wk of Feb 9)  ',fontsize=8,color='#59636F',ha='right')
a1.text(46,promo[-1]*100+0.2,f"{promo[-1]*100:.1f}% (last wk)",ha='right',fontsize=8.5,fontweight='bold',color=RED); a1.text(0,promo[0]*100+0.5,f"{promo[0]*100:.1f}%",fontsize=8.5,color=INK)
a2.plot(x,[v/1000 for v in egsw],color=NAVY,lw=2.2); a2.set_ylabel('Gross / store-wk\n($K, established)',fontsize=8.5); a2.axvline(30.5,color=GRAY,lw=1,ls='--'); a2.set_ylim(9.5,14.5)
lab=['Jul','Sep','Nov','Jan','Mar','May']; pos=[i for i,w in enumerate(weeks) if w[5:7] in ('07','09','11','01','03','05') and int(w[8:])<=7]
a2.set_xticks(pos); a2.set_xticklabels(lab[:len(pos)],fontsize=8.5)
save(fig,'c5_promo')
print(json.dumps({k:V[k] for k in ['cm2_fy','cm2_exit','take_fy','cm2p_fy','cm2p_exit','be_fy']},default=str))
