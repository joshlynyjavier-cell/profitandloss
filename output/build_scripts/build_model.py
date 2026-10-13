"""Build Sola_Kitchen_Financial_Model.xlsx — formula-driven, auditable."""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter as CL
from openpyxl.comments import Comment

SRC, OUT = sys.argv[1], sys.argv[2]
src = openpyxl.load_workbook(SRC)
wb = openpyxl.Workbook()
wb.remove(wb.active)

# ---------- styles ----------
F = "Arial"
f_base = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_title = Font(name=F, size=14, bold=True, color="1F2A44")
f_h2 = Font(name=F, size=11, bold=True, color="1F2A44")
f_hdr = Font(name=F, size=10, bold=True, color="FFFFFF")
f_in = Font(name=F, size=10, color="0000FF")
f_link = Font(name=F, size=10, color="008000")
f_note = Font(name=F, size=9, italic=True, color="555555")
fill_hdr = PatternFill("solid", fgColor="1F2A44")
fill_in = PatternFill("solid", fgColor="FFFF00")
fill_calc = PatternFill("solid", fgColor="EEF2F7")
fill_key = PatternFill("solid", fgColor="E2EFDA")
thin = Side(style="thin", color="BBBBBB")
b_top = Border(top=Side(style="thin", color="333333"))

USD = '$#,##0;($#,##0);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
USDM = '$#,##0.00,,"M";($#,##0.00,,"M");"-"'
PCT = '0.0%;(0.0%);"-"'
PCT2 = '0.00%;(0.00%);"-"'
NUM = '#,##0;(#,##0);"-"'
DEC3 = '0.000;(0.000);"-"'
PTS = '0.0" pts";(0.0" pts");"-"'


def name(nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


def put(ws, cell, val, font=None, fmt=None, fill=None, bold=False, wrap=False, align=None):
    c = ws[cell]
    c.value = val
    c.font = font or (f_bold if bold else f_base)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if wrap or align:
        c.alignment = Alignment(wrap_text=wrap, horizontal=align, vertical="top")
    return c


def hdr(ws, row, col, labels, width=None):
    for i, l in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=l)
        c.font = f_hdr
        c.fill = fill_hdr
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def widths(ws, d):
    for k, v in d.items():
        ws.column_dimensions[k].width = v


def title(ws, t, sub):
    put(ws, "A1", t, f_title)
    put(ws, "A2", sub, f_note)
    ws.sheet_view.showGridLines = False


# =====================================================================
# Raw data sheets (copied value-for-value; calc columns appended to the right)
# =====================================================================
def copy_raw(src_name, dst_name):
    s = src[src_name]
    d = wb.create_sheet(dst_name)
    for row in s.iter_rows():
        for c in row:
            d.cell(row=c.row, column=c.column, value=c.value).font = f_base
    for c in d[1]:
        c.font = f_hdr
        c.fill = fill_hdr
    d.freeze_panes = "A2"
    return d, s.max_row


# placeholder ordering: create front sheets first so tab order reads top-down
front = ["README", "Assumptions", "Exec_Summary", "PL_Waterfall", "Platform_Market",
         "Promo_Escalation", "Incrementality", "Markup", "DoorDash", "Sensitivity", "QA_Checks",
         "Weeks", "Periods"]
for s in front:
    wb.create_sheet(s)

raw_readme, _ = copy_raw("README", "Data_README")
dl, nL = copy_raw("Locations_Master", "Data_Locations")
dp, nP = copy_raw("Marketplace_Payouts", "Data_Payouts")
dc, nC = copy_raw("Channel_Revenue_Summary", "Data_Channel")
dr, nR = copy_raw("Cohort_Retention", "Data_Cohort")
dd_, nD = copy_raw("Daypart_Patterns", "Data_Daypart")
dm, nM = copy_raw("Markup_Test", "Data_Markup")
dt, nT = copy_raw("Contract_Terms", "Data_Contract")
assert nP == 17935 and nL == 123 and nC == 1409 and nR == 2157 and nM == 33

# ---- Data_Payouts calc columns Q:U
for col, lab in zip("QRSTU", ["CALC market_status", "CALC week_idx", "CALC month_key",
                               "CALC total_deductions", "CALC recon_diff (gross-ded-net)"]):
    c = dp[f"{col}1"]; c.value = lab; c.font = f_hdr; c.fill = PatternFill("solid", fgColor="4F6D8F")
for r in range(2, nP + 1):
    dp[f"Q{r}"] = f"=INDEX(Data_Locations!$G$2:$G$123,MATCH(C{r},Data_Locations!$A$2:$A$123,0))"
    dp[f"R{r}"] = f"=MATCH(F{r},Weeks!$A$6:$A$56,0)"
    dp[f"S{r}"] = f"=LEFT(F{r},7)"
    dp[f"T{r}"] = f"=SUM(J{r}:O{r})"
    dp[f"U{r}"] = f"=ROUND(H{r}-T{r}-P{r},2)"
widths(dp, {c: 13 for c in "ABCDEFGHIJKLMNOPQRSTU"})
dp.column_dimensions["D"].width = 26
R_ = f"$2:$"  # helper
rng = lambda col: f"Data_Payouts!${col}$2:${col}${nP}"
for nm, col in [("pPlat", "B"), ("pLoc", "C"), ("pReg", "E"), ("pWeek", "F"), ("pOrders", "G"), ("pGross", "H"),
                ("pTax", "I"), ("pComm", "J"), ("pMkt", "K"), ("pPromo", "L"), ("pAds", "M"), ("pSvc", "N"),
                ("pRef", "O"), ("pNet", "P"), ("pStat", "Q"), ("pWk", "R"), ("pMonth", "S"), ("pDed", "T"),
                ("pDiff", "U")]:
    name(nm, rng(col))

# ---- Data_Channel calc columns M:P
for col, lab in zip("MNOP", ["CALC same_store (12 months)", "CALC month_key", "CALC weeks_in_month",
                             "CALC quarter_bucket"]):
    c = dc[f"{col}1"]; c.value = lab; c.font = f_hdr; c.fill = PatternFill("solid", fgColor="4F6D8F")
for r in range(2, nC + 1):
    dc[f"M{r}"] = f"=COUNTIF($A$2:$A${nC},A{r})=12"
    dc[f"N{r}"] = f"=LEFT(F{r},7)"
    dc[f"O{r}"] = f"=COUNTIF(Weeks!$D$6:$D$56,N{r})"
    dc[f"P{r}"] = f'=IF(N{r}<="2025-09","Q1 Jul-Sep",IF(N{r}>="2026-04","Q4 Apr-Jun","Mid"))'
widths(dc, {c: 14 for c in "ABCDEFGHIJKLMNOP"})
crng = lambda col: f"Data_Channel!${col}$2:${col}${nC}"
for nm, col in [("cLoc", "A"), ("cInStore", "G"), ("cPickup", "H"), ("c1P", "I"), ("c3PD", "J"),
                ("cCater", "K"), ("cTotal", "L"), ("cSame", "M"), ("cMonth", "N"), ("cQ", "P"), ("cStat", "E")]:
    name(nm, crng(col))

# ---- Data_Cohort calc columns L:M
for col, lab in zip("LM", ["CALC retained_30d (n)", "CALC retained_90d (n)"]):
    c = dr[f"{col}1"]; c.value = lab; c.font = f_hdr; c.fill = PatternFill("solid", fgColor="4F6D8F")
for r in range(2, nR + 1):
    dr[f"L{r}"] = f"=H{r}*I{r}"
    dr[f"M{r}"] = f"=H{r}*K{r}"
rrng = lambda col: f"Data_Cohort!${col}$2:${col}${nR}"
for nm, col in [("rReg", "C"), ("rStat", "D"), ("rChan", "F"), ("rPromo", "G"), ("rNew", "H"), ("rR30", "L"),
                ("rR90", "M")]:
    name(nm, rrng(col))
widths(dr, {c: 14 for c in "ABCDEFGHIJKLM"})

# ---- Data_Markup calc columns G:I
for col, lab in zip("GHI", ["CALC net_rev_per_store_idx (vol x net/order)", "CALC ln(volume idx)",
                            "CALC ln(1+markup)"]):
    c = dm[f"{col}1"]; c.value = lab; c.font = f_hdr; c.fill = PatternFill("solid", fgColor="4F6D8F")
for r in range(2, nM + 1):
    dm[f"G{r}"] = f"=D{r}*E{r}"
    dm[f"H{r}"] = f"=LN(D{r})"
    dm[f"I{r}"] = f"=LN(1+C{r}/100)"
widths(dm, {"A": 10, "B": 30, "C": 12, "D": 14, "E": 14, "F": 60, "G": 16, "H": 12, "I": 12})
widths(dd_, {c: 18 for c in "ABCDEF"})
widths(dt, {"A": 10, "B": 62, "C": 12, "D": 16, "E": 44})
widths(dl, {c: 14 for c in "ABCDEFGHI"}); dl.column_dimensions["B"].width = 30
raw_readme.column_dimensions["A"].width = 160

# =====================================================================
# Assumptions
# =====================================================================
A = wb["Assumptions"]
title(A, "Assumptions & Inputs", "Blue/yellow = editable input. Black = formula. Green = link to another sheet. "
      "Source column classifies each item as OBSERVED (data), ASSIGNMENT (given), or ANALYST (my judgment).")
hdr(A, 4, 1, ["Parameter", "Value", "Unit", "Source / classification", "Rationale / note"])
widths(A, {"A": 46, "B": 14, "C": 10, "D": 30, "E": 90})
arow = [5]


def ainput(label, nm, val, unit, srcl, note, fmt=PCT, is_input=True, section=None):
    r = arow[0]
    if section:
        put(A, f"A{r}", section, f_h2); arow[0] += 1; r = arow[0]
    put(A, f"A{r}", label)
    c = put(A, f"B{r}", val, f_in if is_input else f_base, fmt, fill_in if is_input else None)
    put(A, f"C{r}", unit); put(A, f"D{r}", srcl); put(A, f"E{r}", note, wrap=False)
    name(nm, f"Assumptions!$B${r}")
    arow[0] += 1
    return r


ainput("Food & packaging COGS (% of gross subtotal)", "COGS_PCT", 0.30, "%", "ASSIGNMENT",
       "Assignment starting point; applied to un-marked-up gross (see Markup sheet).", section="Cost structure")
ainput("Variable delivery labor (% of gross subtotal)", "LABOR_PCT", 0.16, "%", "ASSIGNMENT",
       "Assignment starting point (dine-in ~22% incl. service).")
ainput("Dine-in contribution benchmark (% of sales)", "DINEIN_CM", 0.48, "%", "ASSIGNMENT",
       "Applied to any sale cannibalized from dine-in / own channels (pickup & 1P assumed similar — ANALYST).")
ainput("Weeks per restaurant year (13 x 28 days / 7)", "WEEKS_PER_YEAR", 52, "weeks", "ASSIGNMENT",
       "13 periods of 28 days = 364 days.", fmt=NUM, section="Period normalization")
ainput("Weeks in payout data", "WEEKS_IN_DATA", "=COUNTA(Weeks!$A$6:$A$56)", "weeks", "OBSERVED",
       "Jul 7 2025 – Jun 22 2026 week-starts; README says 'Jul 2025 – Jun 2026' but only 51 weeks exist.",
       fmt=NUM, is_input=False)
ainput("Annualization factor (52 / weeks in data)", "ANN", "=WEEKS_PER_YEAR/WEEKS_IN_DATA", "x", "DERIVED",
       "FY figures = 51-week actuals x this factor.", fmt="0.0000", is_input=False)
ainput("First week index of latest 28-day period (exit window)", "EXIT_START", "=WEEKS_IN_DATA-3", "idx",
       "DERIVED", "Exit run-rate = latest 4 weeks x 13 periods.", fmt=NUM, is_input=False)
ainput("Periods per year", "PERIODS", 13, "#", "ASSIGNMENT", "", fmt=NUM)
ainput("Last pre-escalation week index (promo baseline window = weeks 1..this)", "PROMO_PRE_END", 31, "idx",
       "ANALYST", "Week of 2026-02-02. Promo % of gross is flat ~1.5% through this week, then ramps every week. "
       "Chosen by inspection of Weeks!R.", fmt=NUM, section="Promo escalation")
ainput("Promo rollback: share of escalation removed", "ROLLBACK_SHARE", 1.0, "%", "ANALYST",
       "Return promo to the pre-Feb-2026 % of gross.")
ainput("Promo rollback: assumed gross volume loss", "ROLLBACK_VOL_LOSS", 0.03, "%", "ANALYST",
       "Gross/store-week did not rise as promo went 1.5%→7.4%, so expected loss is ~0; 3% haircut for prudence.")
ainput("DoorDash current commission", "DD_COMM_CUR", 0.27, "%", "OBSERVED (Contract_Terms)",
       "Verified: commission/gross for DoorDash in payouts = 27.0% (see QA).", section="DoorDash contract")
ainput("DoorDash proposed commission", "DD_COMM_PROP", 0.31, "%", "OBSERVED (Contract_Terms)", "")
ainput("Proposed minimum ads share of DoorDash marketing", "DD_ADS_MIN", 0.60, "%", "OBSERVED (Contract_Terms)",
       "Current DoorDash ad share is derived from payouts (DoorDash sheet).")
ainput("DoorDash gross growth, year 2", "DD_GROWTH", 0.08, "%", "ASSIGNMENT (Contract_Terms note)",
       "'grew ~8% last yr'. Year-1 base = model-derived annualized DoorDash gross.")
ainput("Counter: headline commission", "CTR_COMM", 0.285, "%", "ANALYST",
       "Negotiating position: concede 1.5 pts in exchange for the 24-month volume commitment DoorDash wants.")
ainput("Counter: max ads share (spend-neutral)", "CTR_ADS", 0.40, "%", "ANALYST", "")
ainput("Walk: share of DoorDash gross that switches to UberEats/Grubhub", "WALK_R3", 0.30, "%", "ANALYST",
       "No observed data. Most DoorDash users multi-home; flagged as key uncertainty (see grid on DoorDash sheet).")
ainput("Walk: share of DoorDash gross recaptured in own channels", "WALK_ROWN", 0.15, "%", "ANALYST",
       "Capped low: lunch/dinner seat utilization is 70–95% (Daypart), limiting dine-in recapture at peak.")
ainput("Incrementality — Bear (share of 3PD sales that are truly incremental)", "INC_BEAR", 0.35, "%", "ANALYST",
       "Below observed estimate.", section="Incrementality scenarios")
ainput("Incrementality — Base", "INC_BASE", 0.55, "%", "ANALYST",
       "Rounded from relative-growth estimate on Incrementality sheet (~53%).")
ainput("Incrementality — Bull", "INC_BULL", 0.80, "%", "ANALYST", "Most 3PD orders would not otherwise happen.")
ainput("Menu markup to apply (5, 10 or 15 — tested levels)", "MARKUP_SEL", 5, "%", "ANALYST",
       "5% is the only level where test volume is ~flat; 10–15% add little CM2 for more volume risk.",
       fmt="0", section="Menu markup")
ainput("Haircut on test-implied markup gain", "MARKUP_HAIRCUT", 0.50, "%", "ANALYST",
       "n=8 stores per arm, DoorDash only, 'deliberately noisy' — bank half.")

r = arow[0] + 1
put(A, f"A{r}", "Color legend", f_h2)
put(A, f"A{r+1}", "Blue text on yellow = input", f_in, fill=fill_in)
put(A, f"A{r+2}", "Black = formula on this sheet")
put(A, f"A{r+3}", "Green = link pulled from another sheet", f_link)

# =====================================================================
# Weeks (51 rows)
# =====================================================================
W = wb["Weeks"]
title(W, "Weekly roll-up (from Data_Payouts)", "One row per payout week. All $ via SUMIFS on Data_Payouts. "
      "Period = 28-day restaurant period counted back from the latest week (P1 is a 3-week stub).")
weeks = sorted({src["Marketplace_Payouts"].cell(row=r, column=6).value for r in range(2, nP + 1)})
assert len(weeks) == 51
cols = ["week_start", "week_idx", "period", "month_key", "stores_active", "orders", "gross", "commission",
        "mkt_comm", "promo", "ads", "service_fee", "refunds", "net_deposit", "COGS", "labor", "CM2",
        "promo % gross", "CM2 % gross", "gross / store-wk", "orders / store-wk", "CM2 / store-wk",
        "Estab. stores", "Estab. gross / store-wk", "Estab. orders / store-wk", "DD promo %", "UE promo %",
        "GH promo %"]
hdr(W, 5, 1, cols)
for i, w in enumerate(weeks):
    r = 6 + i
    W[f"A{r}"] = w
    W[f"B{r}"] = f"=ROW()-5"
    W[f"C{r}"] = f"=PERIODS-INT((WEEKS_IN_DATA-B{r})/4)"
    W[f"D{r}"] = f"=LEFT(A{r},7)"
    W[f"E{r}"] = f'=COUNTIFS(pWk,B{r},pPlat,"DoorDash")'
    for col, rn in zip("FGHIJKLMN", ["pOrders", "pGross", "pComm", "pMkt", "pPromo", "pAds", "pSvc", "pRef", "pNet"]):
        W[f"{col}{r}"] = f"=SUMIFS({rn},pWk,$B{r})"
    W[f"O{r}"] = f"=G{r}*COGS_PCT"
    W[f"P{r}"] = f"=G{r}*LABOR_PCT"
    W[f"Q{r}"] = f"=N{r}-O{r}-P{r}"
    W[f"R{r}"] = f"=J{r}/G{r}"
    W[f"S{r}"] = f"=Q{r}/G{r}"
    W[f"T{r}"] = f"=G{r}/E{r}"
    W[f"U{r}"] = f"=F{r}/E{r}"
    W[f"V{r}"] = f"=Q{r}/E{r}"
    W[f"W{r}"] = f'=COUNTIFS(pWk,B{r},pPlat,"DoorDash",pStat,"Established")'
    W[f"X{r}"] = f'=SUMIFS(pGross,pWk,B{r},pStat,"Established")/W{r}'
    W[f"Y{r}"] = f'=SUMIFS(pOrders,pWk,B{r},pStat,"Established")/W{r}'
    W[f"Z{r}"] = f'=SUMIFS(pPromo,pWk,B{r},pPlat,"DoorDash")/SUMIFS(pGross,pWk,B{r},pPlat,"DoorDash")'
    W[f"AA{r}"] = f'=SUMIFS(pPromo,pWk,B{r},pPlat,"UberEats")/SUMIFS(pGross,pWk,B{r},pPlat,"UberEats")'
    W[f"AB{r}"] = f'=SUMIFS(pPromo,pWk,B{r},pPlat,"Grubhub")/SUMIFS(pGross,pWk,B{r},pPlat,"Grubhub")'
    for col in "EFW":
        W[f"{col}{r}"].number_format = NUM
    for col in ["G", "H", "I", "J", "K", "L", "M", "N", "O", "P", "Q", "T", "V", "X"]:
        W[f"{col}{r}"].number_format = USD
    for col in ["U", "Y"]:
        W[f"{col}{r}"].number_format = "#,##0.0"
    for col in ["R", "S", "Z", "AA", "AB"]:
        W[f"{col}{r}"].number_format = PCT2
    for c in W[r]:
        c.font = f_base
r = 57
put(W, f"A{r}", "TOTAL", bold=True)
for col in "FGHIJKLMNOPQ":
    W[f"{col}{r}"] = f"=SUM({col}6:{col}56)"; W[f"{col}{r}"].number_format = USD if col != "F" else NUM
    W[f"{col}{r}"].font = f_bold; W[f"{col}{r}"].border = b_top
W.freeze_panes = "B6"
widths(W, {CL(i): 12 for i in range(1, 29)})

# =====================================================================
# Periods (13 x 28-day)
# =====================================================================
Pd = wb["Periods"]
title(Pd, "28-day restaurant periods (13 per year)", "Built from Weeks. P1 = 3-week stub (Jul 7–21 2025); "
      "P13 = latest full period (Jun 1–28 2026) = exit window.")
hdr(Pd, 5, 1, ["Period", "Weeks", "Store-weeks", "Gross", "Promo", "Net deposit", "CM2", "Promo % gross",
               "Take rate", "CM2 % gross", "Gross / store-wk", "CM2 / store-wk", "Estab. gross / store-wk"])
for i in range(13):
    r = 6 + i
    Pd[f"A{r}"] = i + 1
    Pd[f"B{r}"] = f"=COUNTIF(Weeks!$C$6:$C$56,A{r})"
    Pd[f"C{r}"] = f"=SUMIF(Weeks!$C$6:$C$56,A{r},Weeks!$E$6:$E$56)"
    Pd[f"D{r}"] = f"=SUMIF(Weeks!$C$6:$C$56,A{r},Weeks!$G$6:$G$56)"
    Pd[f"E{r}"] = f"=SUMIF(Weeks!$C$6:$C$56,A{r},Weeks!$J$6:$J$56)"
    Pd[f"F{r}"] = f"=SUMIF(Weeks!$C$6:$C$56,A{r},Weeks!$N$6:$N$56)"
    Pd[f"G{r}"] = f"=SUMIF(Weeks!$C$6:$C$56,A{r},Weeks!$Q$6:$Q$56)"
    Pd[f"H{r}"] = f"=E{r}/D{r}"
    Pd[f"I{r}"] = f"=1-F{r}/D{r}"
    Pd[f"J{r}"] = f"=G{r}/D{r}"
    Pd[f"K{r}"] = f"=D{r}/C{r}"
    Pd[f"L{r}"] = f"=G{r}/C{r}"
    Pd[f"M{r}"] = (f"=SUMPRODUCT((Weeks!$C$6:$C$56=A{r})*Weeks!$X$6:$X$56*Weeks!$W$6:$W$56)"
                   f"/SUMIF(Weeks!$C$6:$C$56,A{r},Weeks!$W$6:$W$56)")
    for c, fm in zip("ABCDEFGHIJKLM", [NUM, NUM, NUM, USD, USD, USD, USD, PCT, PCT, PCT, USD, USD, USD]):
        Pd[f"{c}{r}"].number_format = fm; Pd[f"{c}{r}"].font = f_base
Pd["A19"] = "Total"; Pd["A19"].font = f_bold
for c in "BCDEFG":
    Pd[f"{c}19"] = f"=SUM({c}6:{c}18)"; Pd[f"{c}19"].number_format = USD if c > "C" else NUM; Pd[f"{c}19"].font = f_bold
widths(Pd, {CL(i): 14 for i in range(1, 14)})

# =====================================================================
# PL_Waterfall
# =====================================================================
PL = wb["PL_Waterfall"]
title(PL, "Marketplace contribution P&L — gross to CM2 (all platforms, all stores)",
      "Built bottom-up from 17,934 weekly x store x platform payout rows. Deductions shown negative. "
      "Tax collected/remitted by platform is excluded (memo only).")
hdr(PL, 4, 1, ["Line", "51-wk actual ($)", "% of gross", "$ / order", "FY annualized (52 wk)",
               "Exit run-rate (latest period x13)", "Exit % of gross", "Exit vs FY (pts)", "Definition / source"])
widths(PL, {"A": 40, "B": 16, "C": 11, "D": 11, "E": 18, "F": 20, "G": 12, "H": 12, "I": 70})
lines = [
    ("Orders", "=SUM(pOrders)", f'=SUMIFS(pOrders,pWk,">="&EXIT_START)', "n_orders", NUM),
    ("Gross subtotal", "=SUM(pGross)", f'=SUMIFS(pGross,pWk,">="&EXIT_START)', "Menu price paid, pre-deduction", USD),
    ("Less: platform commission", "=-SUM(pComm)", f'=-SUMIFS(pComm,pWk,">="&EXIT_START)', "commission", USD),
    ("Less: marketing commission", "=-SUM(pMkt)", f'=-SUMIFS(pMkt,pWk,">="&EXIT_START)', "marketing_commission", USD),
    ("Less: restaurant-funded promotions", "=-SUM(pPromo)", f'=-SUMIFS(pPromo,pWk,">="&EXIT_START)', "promo_funded", USD),
    ("Less: advertising", "=-SUM(pAds)", f'=-SUMIFS(pAds,pWk,">="&EXIT_START)', "ad_spend", USD),
    ("Less: service fees", "=-SUM(pSvc)", f'=-SUMIFS(pSvc,pWk,">="&EXIT_START)', "service_fee", USD),
    ("Less: refunds & adjustments", "=-SUM(pRef)", f'=-SUMIFS(pRef,pWk,">="&EXIT_START)', "refunds", USD),
]
r = 5
for lab, f1, f2, d, fm in lines:
    put(PL, f"A{r}", lab, bold=(lab in ("Gross subtotal",)))
    PL[f"B{r}"] = f1; PL[f"F{r}"] = f2; put(PL, f"I{r}", d, f_note)
    r += 1
# r now 13 -> Net
put(PL, "A13", "Net revenue / bank deposit", bold=True)
PL["B13"] = "=SUM(B6:B12)"; PL["F13"] = "=SUM(F6:F12)"
put(PL, "I13", "Gross less all six platform deductions (README definition). Ties to source net_bank_deposit (row 25).", f_note)
put(PL, "A14", "Effective all-in take rate", bold=True)
PL["B14"] = "=1-B13/B6"; PL["F14"] = "=1-F13/F6"
put(PL, "I14", "(Gross − net) / gross; includes promo & ads, not just headline commission", f_note)
put(PL, "A15", "Less: food & packaging COGS")
PL["B15"] = "=-B6*COGS_PCT"; PL["F15"] = "=-F6*COGS_PCT"
put(PL, "I15", "Assumptions!COGS_PCT x gross", f_note)
put(PL, "A16", "CM1", bold=True); PL["B16"] = "=B13+B15"; PL["F16"] = "=F13+F15"
put(PL, "A17", "Less: variable delivery labor"); PL["B17"] = "=-B6*LABOR_PCT"; PL["F17"] = "=-F6*LABOR_PCT"
put(PL, "I17", "Assumptions!LABOR_PCT x gross", f_note)
put(PL, "A18", "CM2 / contribution margin", bold=True); PL["B18"] = "=B16+B17"; PL["F18"] = "=F16+F17"
put(PL, "A20", "Dine-in contribution benchmark", bold=True); PL["C20"] = "=DINEIN_CM"; PL["G20"] = "=DINEIN_CM"
put(PL, "A21", "CM2 gap vs dine-in (pts of sales)"); PL["C21"] = "=(C18-C20)*100"; PL["G21"] = "=(G18-G20)*100"
put(PL, "A22", "3PD keeps this share of dine-in contribution per $"); PL["C22"] = "=C18/C20"; PL["G22"] = "=G18/G20"
put(PL, "A23", "Break-even incrementality (1 − CM2%/dine-in CM%)"); PL["C23"] = "=1-C18/C20"; PL["G23"] = "=1-G18/G20"
put(PL, "I23", "Share of 3PD sales that must be truly incremental for the channel to add value vs dine-in", f_note)
put(PL, "A25", "Memo: net_bank_deposit per source"); PL["B25"] = "=SUM(pNet)"; PL["F25"] = '=SUMIFS(pNet,pWk,">="&EXIT_START)'
put(PL, "A26", "Memo: check (row 13 − row 25)"); PL["B26"] = "=ROUND(B13-B25,0)"; PL["F26"] = "=ROUND(F13-F25,0)"
put(PL, "A27", "Memo: tax collected & remitted (excluded)"); PL["B27"] = "=SUM(pTax)"
put(PL, "I27", "Marketplace-facilitator tax — never restaurant revenue; not deducted, not added.", f_note)
put(PL, "A28", "Memo: restaurant marketing spend (promo + ads)"); PL["B28"] = "=-(B9+B10)"; PL["F28"] = "=-(F9+F10)"
put(PL, "A29", "Memo: ads share of marketing"); PL["B29"] = "=-B10/B28"; PL["F29"] = "=-F10/F28"
for rr in range(5, 13):
    PL[f"F{rr}"] = f"=({PL[f'F{rr}'].value[1:]})*PERIODS"
for rr in list(range(5, 19)) + [25, 27, 28]:
    if rr == 14:
        continue
    if rr != 5:
        PL[f"C{rr}"] = f"=B{rr}/$B$6"
        PL[f"D{rr}"] = f"=B{rr}/$B$5"
    PL[f"E{rr}"] = f"=B{rr}*ANN"
    if rr <= 18 and rr != 5:
        PL[f"G{rr}"] = f"=F{rr}/$F$6"
        PL[f"H{rr}"] = f"=(G{rr}-C{rr})*100"
PL["C14"] = "=B14"; PL["G14"] = "=F14"; PL["H14"] = "=(G14-C14)*100"
# Exit memo rows need x13 too
PL["F25"] = '=SUMIFS(pNet,pWk,">="&EXIT_START)*PERIODS'
PL["F28"] = "=-(F9+F10)"
for rr in range(5, 30):
    for col in "BEF":
        if PL[f"{col}{rr}"].value is not None:
            PL[f"{col}{rr}"].number_format = NUM if rr == 5 else USD
    for col in "CG":
        if PL[f"{col}{rr}"].value is not None:
            PL[f"{col}{rr}"].number_format = PCT
    if PL[f"D{rr}"].value is not None:
        PL[f"D{rr}"].number_format = USD2
    if PL[f"H{rr}"].value is not None:
        PL[f"H{rr}"].number_format = PTS
    for col in "BCDEFGH":
        PL[f"{col}{rr}"].font = f_bold if rr in (6, 13, 14, 16, 18) else f_base
for rr in (13, 16, 18):
    for col in "ABCDEFGH":
        PL[f"{col}{rr}"].border = b_top
        PL[f"{col}{rr}"].fill = fill_key if rr == 18 else fill_calc
PL["C21"].number_format = PTS; PL["G21"].number_format = PTS
PL["B26"].number_format = NUM; PL["F26"].number_format = NUM
PL["B29"].number_format = PCT; PL["F29"].number_format = PCT
PL.freeze_panes = "B5"

# =====================================================================
# Platform_Market — segment P&Ls
# =====================================================================
PM = wb["Platform_Market"]
title(PM, "Contribution by platform, market (region) and market status",
      "Each column = SUMIFS on Data_Payouts with the filters shown in rows 5–7. Exit columns use latest 28-day period only.")
segs = [("Total", None, None, None)]
segs += [(p, p, None, None) for p in ["DoorDash", "UberEats", "Grubhub"]]
regions = ["Texas", "Southeast", "Southern California", "Mountain West", "Pacific Northwest", "Northeast"]
segs += [(g, None, g, None) for g in regions]
segs += [("Established", None, None, "Established"), ("East Coast (New)", None, None, "New (East Coast launch)")]
segs += [(f"{p[:2] if p!='UberEats' else 'UE'} · {s[:5]}", p, None, st) for p in ["DoorDash", "UberEats", "Grubhub"]
         for s, st in [("Estab", "Established"), ("New", "New (East Coast launch)")]]
segs = [(l.replace("Do ·", "DD ·").replace("Gr ·", "GH ·"), a, b, c) for l, a, b, c in segs]
put(PM, "A5", "Filter: platform", f_note); put(PM, "A6", "Filter: region", f_note); put(PM, "A7", "Filter: market status", f_note)
put(PM, "A4", "Segment", bold=True)
groups = {2: "ALL", 3: "PLATFORM", 6: "REGION / MARKET", 12: "MARKET STATUS", 14: "PLATFORM x STATUS"}
for i, (lab, p, g, s) in enumerate(segs):
    col = CL(2 + i)
    c = PM[f"{col}4"]; c.value = lab; c.font = f_hdr; c.fill = fill_hdr
    c.alignment = Alignment(horizontal="center", wrap_text=True)
    PM[f"{col}5"] = p or "(all)"; PM[f"{col}6"] = g or "(all)"; PM[f"{col}7"] = s or "(all)"
    for rr in (5, 6, 7):
        PM[f"{col}{rr}"].font = f_note; PM[f"{col}{rr}"].alignment = Alignment(horizontal="center", wrap_text=True)
for k, v in groups.items():
    put(PM, f"{CL(k)}3", v, f_note)
lines_pm = [
    ("Orders", "pOrders", NUM), ("Gross subtotal", "pGross", USD), ("Commission", "pComm", USD),
    ("Marketing commission", "pMkt", USD), ("Restaurant-funded promo", "pPromo", USD), ("Advertising", "pAds", USD),
    ("Service fees", "pSvc", USD), ("Refunds & adjustments", "pRef", USD),
]
ROW = {}
r = 9
for lab, rn, fm in lines_pm:
    ROW[lab] = r; put(PM, f"A{r}", lab, bold=(lab == "Gross subtotal"))
    for i, (sl, p, g, s) in enumerate(segs):
        col = CL(2 + i)
        crit = ""
        if p: crit += f",pPlat,{col}$5"
        if g: crit += f",pReg,{col}$6"
        if s: crit += f",pStat,{col}$7"
        PM[f"{col}{r}"] = f"=SUM({rn})" if not crit else f"=SUMIFS({rn}{crit})"
        PM[f"{col}{r}"].number_format = fm
    r += 1
# derived rows
der = [
    ("Net revenue / deposit", "={c}10-SUM({c}11:{c}16)", USD, True),
    ("All-in take rate", "=1-{c}17/{c}10", PCT, False),
    ("  of which commission %", "={c}11/{c}10", PCT, False),
    ("  of which promo + ads %", "=({c}13+{c}14)/{c}10", PCT, False),
    ("Food & packaging COGS", "={c}10*COGS_PCT", USD, False),
    ("CM1", "={c}17-{c}21", USD, True),
    ("CM1 %", "={c}22/{c}10", PCT, False),
    ("Variable delivery labor", "={c}10*LABOR_PCT", USD, False),
    ("CM2 / contribution", "={c}22-{c}24", USD, True),
    ("CM2 % of gross", "={c}25/{c}10", PCT, True),
    ("Average order value (gross)", "={c}10/{c}9", USD2, False),
    ("CM2 per order", "={c}25/{c}9", USD2, True),
    ("Revenue mix (share of total gross)", "={c}10/$B$10", PCT, False),
    ("Contribution mix (share of total CM2)", "={c}25/$B$25", PCT, False),
    ("CM2 % as share of dine-in 48%", "={c}26/DINEIN_CM", PCT, False),
    ("Break-even incrementality", "=1-{c}26/DINEIN_CM", PCT, True),
    ("FY annualized gross", "={c}10*ANN", USD, False),
    ("FY annualized CM2", "={c}25*ANN", USD, True),
    ("Exit period: gross (latest 4 wks)", None, USD, False),
    ("Exit period: promo % of gross", None, PCT, False),
    ("Exit period: CM2 % of gross", None, PCT, True),
    ("Exit run-rate CM2 (x13)", None, USD, True),
]
for lab, fml, fm, b in der:
    ROW[lab] = r; put(PM, f"A{r}", lab, bold=b)
    for i, (sl, p, g, s) in enumerate(segs):
        col = CL(2 + i)
        if fml:
            PM[f"{col}{r}"] = fml.format(c=col)
        else:
            crit = f',pWk,">="&EXIT_START'
            if p: crit += f",pPlat,{col}$5"
            if g: crit += f",pReg,{col}$6"
            if s: crit += f",pStat,{col}$7"
            if lab.startswith("Exit period: gross"):
                PM[f"{col}{r}"] = f"=SUMIFS(pGross{crit})"
            elif "promo" in lab:
                PM[f"{col}{r}"] = f"=SUMIFS(pPromo{crit})/{col}{r-1}"
            elif "CM2 %" in lab:
                PM[f"{col}{r}"] = f"=(SUMIFS(pNet{crit})-{col}{r-2}*(COGS_PCT+LABOR_PCT))/{col}{r-2}"
            else:
                PM[f"{col}{r}"] = f"={col}{r-1}*{col}{r-3}*PERIODS"
        PM[f"{col}{r}"].number_format = fm
        PM[f"{col}{r}"].font = f_bold if b else f_base
        if lab in ("CM2 / contribution", "CM2 % of gross"):
            PM[f"{col}{r}"].fill = fill_key
    r += 1
for rr in range(9, r):
    for i in range(len(segs) + 1):
        c = PM.cell(row=rr, column=1 + i)
        if c.font != f_bold:
            pass
# reconciliation block
r += 1
put(PM, f"A{r}", "Reconciliation checks (should all be 0)", f_h2); r += 1
chk = [("Platforms sum to total — gross", "=ROUND(SUM(C10:E10)-B10,2)"),
       ("Platforms sum to total — CM2", "=ROUND(SUM(C25:E25)-B25,2)"),
       ("Regions sum to total — gross", "=ROUND(SUM(F10:K10)-B10,2)"),
       ("Regions sum to total — CM2", "=ROUND(SUM(F25:K25)-B25,2)"),
       ("Status sum to total — CM2", "=ROUND(SUM(L25:M25)-B25,2)"),
       ("Platform x status sum to total — CM2", "=ROUND(SUM(N25:S25)-B25,2)"),
       ("Total ties to PL_Waterfall CM2", "=ROUND(B25-PL_Waterfall!B18,2)"),
       ("Net (computed) ties to source net deposit", "=ROUND(B17-SUM(pNet),0)")]
PMCHK = r
for lab, f in chk:
    put(PM, f"A{r}", lab); PM[f"B{r}"] = f; PM[f"B{r}"].number_format = '#,##0.00'; r += 1
widths(PM, {"A": 38, **{CL(i): 13 for i in range(2, 2 + len(segs))}})
PM.freeze_panes = "B9"
for rr in range(9, PMCHK):
    for i in range(len(segs)):
        c = PM.cell(row=rr, column=2 + i)
        if c.font is None or c.font.name != F:
            c.font = f_base
for rr in range(9, 17):
    for i in range(len(segs)):
        PM.cell(row=rr, column=2 + i).font = f_base

# =====================================================================
# Promo_Escalation
# =====================================================================
PE = wb["Promo_Escalation"]
title(PE, "Restaurant-funded promo escalated ~5x from Feb-2026 with no volume response",
      "Pre window = weeks 1..PROMO_PRE_END (Jul 7 2025 – Feb 2 2026). Exit = latest 28-day period. "
      "Established stores used for volume to remove East Coast opening mix.")
hdr(PE, 4, 1, ["Metric", "Baseline (wks 1–31 unless noted)", "Exit period (latest 4 wks)", "Change", "Note"])
widths(PE, {"A": 52, "B": 18, "C": 18, "D": 14, "E": 70})
pre = lambda rn: f'SUMIFS({rn},pWk,"<="&PROMO_PRE_END)'
ex = lambda rn: f'SUMIFS({rn},pWk,">="&EXIT_START)'
pre_e = lambda rn: f'SUMIFS({rn},pWk,"<="&PROMO_PRE_END,pStat,"Established")'
ex_e = lambda rn: f'SUMIFS({rn},pWk,">="&EXIT_START,pStat,"Established")'
sw_pre = 'SUMIFS(Weeks!$W$6:$W$56,Weeks!$B$6:$B$56,"<="&PROMO_PRE_END)'
sw_ex = 'SUMIFS(Weeks!$W$6:$W$56,Weeks!$B$6:$B$56,">="&EXIT_START)'
l4 = lambda rn: (f'SUMIFS({rn},pWk,">="&(PROMO_PRE_END-3),pWk,"<="&PROMO_PRE_END,pStat,"Established")'
                 '/SUMIFS(Weeks!$W$6:$W$56,Weeks!$B$6:$B$56,">="&(PROMO_PRE_END-3),Weeks!$B$6:$B$56,"<="&PROMO_PRE_END)')
pe_rows = [
    ("Promo % of gross — all stores", f"={pre('pPromo')}/{pre('pGross')}", f"={ex('pPromo')}/{ex('pGross')}", PCT2, "pts"),
    ("Promo % of gross — DoorDash", f'=SUMIFS(pPromo,pWk,"<="&PROMO_PRE_END,pPlat,"DoorDash")/SUMIFS(pGross,pWk,"<="&PROMO_PRE_END,pPlat,"DoorDash")',
     f'=SUMIFS(pPromo,pWk,">="&EXIT_START,pPlat,"DoorDash")/SUMIFS(pGross,pWk,">="&EXIT_START,pPlat,"DoorDash")', PCT2, "pts"),
    ("Promo % of gross — UberEats", f'=SUMIFS(pPromo,pWk,"<="&PROMO_PRE_END,pPlat,"UberEats")/SUMIFS(pGross,pWk,"<="&PROMO_PRE_END,pPlat,"UberEats")',
     f'=SUMIFS(pPromo,pWk,">="&EXIT_START,pPlat,"UberEats")/SUMIFS(pGross,pWk,">="&EXIT_START,pPlat,"UberEats")', PCT2, "pts"),
    ("Promo % of gross — East Coast (New)", f'=SUMIFS(pPromo,pWk,"<="&PROMO_PRE_END,pStat,"New (East Coast launch)")/SUMIFS(pGross,pWk,"<="&PROMO_PRE_END,pStat,"New (East Coast launch)")',
     f'=SUMIFS(pPromo,pWk,">="&EXIT_START,pStat,"New (East Coast launch)")/SUMIFS(pGross,pWk,">="&EXIT_START,pStat,"New (East Coast launch)")', PCT2, "pts"),
    ("Established: gross per store-week (last 4 baseline wks → exit)", "=" + l4('pGross'), f"={ex_e('pGross')}/{sw_ex}", USD, "%"),
    ("Established: orders per store-week (last 4 baseline wks → exit)", "=" + l4('pOrders'), f"={ex_e('pOrders')}/{sw_ex}", "#,##0.0", "%"),
    ("Memo: Established gross / store-wk, full baseline avg (incl. Jul–Jan growth)", f"={pre_e('pGross')}/{sw_pre}", f"={ex_e('pGross')}/{sw_ex}", USD, "%"),
    ("CM2 % of gross — all stores", f"=({pre('pNet')}-{pre('pGross')}*(COGS_PCT+LABOR_PCT))/{pre('pGross')}",
     f"=({ex('pNet')}-{ex('pGross')}*(COGS_PCT+LABOR_PCT))/{ex('pGross')}", PCT, "pts"),
]
for i, (lab, f1, f2, fm, chg) in enumerate(pe_rows):
    rr = 5 + i
    put(PE, f"A{rr}", lab); PE[f"B{rr}"] = f1; PE[f"C{rr}"] = f2
    PE[f"D{rr}"] = f"=(C{rr}-B{rr})*100" if chg == "pts" else f"=C{rr}/B{rr}-1"
    for col in "BC":
        PE[f"{col}{rr}"].number_format = fm; PE[f"{col}{rr}"].font = f_base
    PE[f"D{rr}"].number_format = PTS if chg == "pts" else PCT
put(PE, "E9", "Last 4 baseline weeks (Jan 12 – Feb 2) vs exit (Jun 1–28): promo ~4.5x, gross/store-wk slightly DOWN. Dine-in also softened seasonally Feb→Jun (Incrementality rows 13–17).", f_note)
put(PE, "A15", "Relative-to-dine-in check (same-store, per store-week; from Incrementality sheet)", f_h2)
put(PE, "A16", "3PD ÷ in-store ratio: monthly growth Jul→Jan (pre)")
PE["B16"] = "=(Incrementality!K12/Incrementality!K6)^(1/6)-1"
put(PE, "A17", "3PD ÷ in-store ratio: monthly growth Feb→Jun (escalation)")
PE["B17"] = "=(Incrementality!K17/Incrementality!K13)^(1/4)-1"
put(PE, "E16", "If promo bought share, the 3PD/dine-in ratio should accelerate after Feb. It decelerated.", f_note)
PE["B16"].number_format = PCT2; PE["B17"].number_format = PCT2
put(PE, "A19", "Cost of the escalation", f_h2)
put(PE, "A20", "Exit annualized gross (latest period x13)"); PE["B20"] = "=PL_Waterfall!F6"; PE["B20"].font = f_link
put(PE, "A21", "Incremental promo vs baseline (pts of gross)"); PE["B21"] = "=C5-B5"
put(PE, "A22", "Annualized cost of escalation at exit run-rate"); PE["B22"] = "=B20*B21"
put(PE, "A23", "Rollback value: promo saved"); PE["B23"] = "=B22*ROLLBACK_SHARE*(1-ROLLBACK_VOL_LOSS)"
put(PE, "A24", "Rollback value: CM2 lost on assumed volume loss")
PE["B24"] = "=-B20*ROLLBACK_VOL_LOSS*PL_Waterfall!G18"
put(PE, "A25", "Net annual CM2 from rollback", bold=True); PE["B25"] = "=B23+B24"
put(PE, "E25", "Uses exit-period CM2% for the lost-volume margin (conservative-ish: lost orders carry exit economics).", f_note)
put(PE, "A26", "Memo: DoorDash commission ask (+4 pts), annual, for scale"); PE["B26"] = "=DoorDash!C12"; PE["B26"].font = f_link
put(PE, "A27", "Escalation cost ÷ DoorDash ask"); PE["B27"] = "=B22/B26"
for rr in (20, 22, 23, 24, 25, 26):
    PE[f"B{rr}"].number_format = USD
PE["B21"].number_format = PCT2; PE["B27"].number_format = "0.0x"
PE["B25"].fill = fill_key; PE["B25"].font = f_bold

# =====================================================================
# Incrementality
# =====================================================================
IN = wb["Incrementality"]
title(IN, "Incrementality & channel trade-offs",
      "A. Same-store channel trends per store-week (Channel_Revenue_Summary; months are 4- or 5-week buckets, so "
      "normalized per week).  B. Implied displacement.  C. Daypart capacity.  D. Cohort retention.  E. Scenarios.")
widths(IN, {"A": 40, **{CL(i): 13 for i in range(2, 16)}})
put(IN, "A4", "A. Same-store (stores with all 12 months) — revenue per store-week", f_h2)
hdr(IN, 5, 1, ["Month", "Weeks", "Stores", "In-store", "Pickup", "1P delivery", "3PD gross", "Catering", "Total",
               "3PD share", "3PD ÷ in-store"])
months = ["2025-07", "2025-08", "2025-09", "2025-10", "2025-11", "2025-12", "2026-01", "2026-02", "2026-03",
          "2026-04", "2026-05", "2026-06"]
for i, m in enumerate(months):
    rr = 6 + i
    IN[f"A{rr}"] = m
    IN[f"B{rr}"] = f"=COUNTIF(Weeks!$D$6:$D$56,A{rr})"
    IN[f"C{rr}"] = f"=COUNTIFS(cMonth,A{rr},cSame,TRUE)"
    for col, rn in zip("DEFGHI", ["cInStore", "cPickup", "c1P", "c3PD", "cCater", "cTotal"]):
        IN[f"{col}{rr}"] = f"=SUMIFS({rn},cMonth,$A{rr},cSame,TRUE)/($B{rr}*$C{rr})"
        IN[f"{col}{rr}"].number_format = USD
    IN[f"J{rr}"] = f"=G{rr}/I{rr}"; IN[f"J{rr}"].number_format = PCT
    IN[f"K{rr}"] = f"=G{rr}/D{rr}"; IN[f"K{rr}"].number_format = "0.000"
# quarter compare
put(IN, "A19", "Q1 (Jul–Sep, 13 wks) per store-week")
put(IN, "A20", "Q4 (Apr–Jun, 12 wks) per store-week")
put(IN, "A21", "Growth Q1→Q4", bold=True)
for col in "DEFGHI":
    IN[f"{col}19"] = f"=SUMPRODUCT({col}6:{col}8,$B$6:$B$8*$C$6:$C$8)/SUMPRODUCT($B$6:$B$8,$C$6:$C$8)"
    IN[f"{col}20"] = f"=SUMPRODUCT({col}15:{col}17,$B$15:$B$17*$C$15:$C$17)/SUMPRODUCT($B$15:$B$17,$C$15:$C$17)"
    IN[f"{col}21"] = f"={col}20/{col}19-1"
    IN[f"{col}19"].number_format = USD; IN[f"{col}20"].number_format = USD; IN[f"{col}21"].number_format = PCT
    IN[f"{col}21"].font = f_bold
IN["J19"] = "=G19/I19"; IN["J20"] = "=G20/I20"; IN["J21"] = "=(J20-J19)*100"
IN["J19"].number_format = PCT; IN["J20"].number_format = PCT; IN["J21"].number_format = PTS
put(IN, "A23", "B. Implied displacement of dine-in by 3PD growth (same-store, Q1→Q4)", f_h2)
put(IN, "A24", "Control growth: pickup + 1P + catering combined")
IN["B24"] = "=(SUM(E20,F20,H20)/SUM(E19,F19,H19))-1"
put(IN, "A25", "In-store growth"); IN["B25"] = "=D21"
put(IN, "A26", "3PD growth"); IN["B26"] = "=G21"
put(IN, "A27", "In-store shortfall vs control ($ / store-wk)"); IN["B27"] = "=(B24-B25)*D19"
put(IN, "A28", "3PD growth above control ($ / store-wk)"); IN["B28"] = "=(B26-B24)*G19"
put(IN, "A29", "Implied displacement ratio (shortfall ÷ excess)", bold=True); IN["B29"] = "=B27/B28"
put(IN, "A30", "Implied incrementality of marginal 3PD growth", bold=True); IN["B30"] = "=1-B29"
put(IN, "A31", "Break-even incrementality (FY, from PL_Waterfall)"); IN["B31"] = "=PL_Waterfall!C23"; IN["B31"].font = f_link
put(IN, "A32", "Cross-check: store-level correl(3PD growth, in-store growth)"); IN["B32"] = "=CORREL(StoreGrowth!F3:F114,StoreGrowth!E3:E114)"
put(IN, "C24", "ASSUMPTION: absent 3PD, in-store would have grown like the other own channels (same seasonality). "
    "One year of data — cannot separate seasonality from trend; result is indicative, not causal.", f_note)
put(IN, "C32", "Positive store-level correlation = no cross-sectional evidence of cannibalization (stores with more 3PD growth "
    "also grew dine-in). Evidence is mixed → treat incrementality as the key uncertainty.", f_note)
for rr, fm in zip(range(24, 33), [PCT, PCT, PCT, USD, USD, PCT, PCT, PCT, "0.00"]):
    IN[f"B{rr}"].number_format = fm
IN["B29"].fill = fill_key; IN["B30"].fill = fill_key

# Daypart
put(IN, "A35", "C. Daypart capacity — share of 3PD orders by daypart vs dine-in seat utilization", f_h2)
hdr(IN, 36, 1, ["Region", "DD orders", "UE orders", "GH orders", "Peak share DD", "Peak share UE", "Peak share GH",
                "3PD orders in peak", "Peak seat util (avg L+D)", "Off-peak seat util (avg A+L)",
                "Util. faced by avg 3PD order"])
for i, g in enumerate(regions):
    rr = 37 + i
    base = 2 + i * 4  # Data_Daypart row of Lunch for this region
    IN[f"A{rr}"] = g
    IN[f"B{rr}"] = f'=SUMIFS(pOrders,pReg,$A{rr},pPlat,"DoorDash")'
    IN[f"C{rr}"] = f'=SUMIFS(pOrders,pReg,$A{rr},pPlat,"UberEats")'
    IN[f"D{rr}"] = f'=SUMIFS(pOrders,pReg,$A{rr},pPlat,"Grubhub")'
    # find rows by region via SUMIFS on daypart table
    dpr = lambda col, part: f'SUMIFS(Data_Daypart!${col}$2:${col}$25,Data_Daypart!$A$2:$A$25,$A{rr},Data_Daypart!$B$2:$B$25,"{part}")'
    tot = lambda col: f'SUMIFS(Data_Daypart!${col}$2:${col}$25,Data_Daypart!$A$2:$A$25,$A{rr})'
    IN[f"E{rr}"] = f"=({dpr('C','Lunch 11-2')}+{dpr('C','Dinner 5-9')})/{tot('C')}"
    IN[f"F{rr}"] = f"=({dpr('D','Lunch 11-2')}+{dpr('D','Dinner 5-9')})/{tot('D')}"
    IN[f"G{rr}"] = f"=({dpr('E','Lunch 11-2')}+{dpr('E','Dinner 5-9')})/{tot('E')}"
    IN[f"H{rr}"] = f"=(B{rr}*E{rr}+C{rr}*F{rr}+D{rr}*G{rr})/SUM(B{rr}:D{rr})"
    IN[f"I{rr}"] = f"=({dpr('F','Lunch 11-2')}+{dpr('F','Dinner 5-9')})/2"
    IN[f"J{rr}"] = f"=({dpr('F','Afternoon 2-5')}+{dpr('F','Late 9-12')})/2"
    parts = ["Lunch 11-2", "Afternoon 2-5", "Dinner 5-9", "Late 9-12"]
    w = "+".join([f"(B{rr}*{dpr('C',p)}/{tot('C')}+C{rr}*{dpr('D',p)}/{tot('D')}+D{rr}*{dpr('E',p)}/{tot('E')})*{dpr('F',p)}" for p in parts])
    IN[f"K{rr}"] = f"=({w})/SUM(B{rr}:D{rr})"
    for col in "BCD":
        IN[f"{col}{rr}"].number_format = NUM
    for col in "EFGHIJK":
        IN[f"{col}{rr}"].number_format = PCT
rr = 43
put(IN, f"A{rr}", "Chain (order-weighted)", bold=True)
for col in "BCD":
    IN[f"{col}{rr}"] = f"=SUM({col}37:{col}42)"; IN[f"{col}{rr}"].number_format = NUM
for col in "EFG":
    oc = {"E": "B", "F": "C", "G": "D"}[col]
    IN[f"{col}{rr}"] = f"=SUMPRODUCT({col}37:{col}42,{oc}37:{oc}42)/{oc}43"; IN[f"{col}{rr}"].number_format = PCT
IN["H43"] = "=SUMPRODUCT(H37:H42,B37:B42+C37:C42+D37:D42)/SUM(B43:D43)"
IN["K43"] = "=SUMPRODUCT(K37:K42,B37:B42+C37:C42+D37:D42)/SUM(B43:D43)"
IN["I43"] = "=AVERAGE(I37:I42)"; IN["J43"] = "=AVERAGE(J37:J42)"
for col in "HIJK":
    IN[f"{col}43"].number_format = PCT; IN[f"{col}43"].font = f_bold
put(IN, "A45", "Data issue: raw daypart order shares sum to 0.92–1.03 per region/platform (not 1.00). Shares are normalized to their region/platform total here.", f_note)
put(IN, "A44", "Reading: ~70% of 3PD orders land at lunch/dinner when 70–95% of seats are already full. Those orders "
    "compete with dine-in for line capacity (README: high util = delivery competes). Off-peak (afternoon/late, ~22–42% "
    "util) is where 3PD is most plausibly incremental.", f_note)

# Cohort
put(IN, "A47", "D. Cohort retention (new-customer weighted) — Nov-2025 to May-2026 cohorts, 44 stores", f_h2)
hdr(IN, 48, 1, ["Channel", "Via promo?", "New customers", "30-day retention", "90-day retention",
                "90d vs dine-in"])
coh = [("Dine-In", False), ("1st Party Delivery", False), ("DoorDash", False), ("DoorDash", True),
       ("UberEats", False), ("UberEats", True), ("Grubhub", False)]
for i, (ch, pr) in enumerate(coh):
    rr = 49 + i
    IN[f"A{rr}"] = ch; IN[f"B{rr}"] = pr
    IN[f"C{rr}"] = f"=SUMIFS(rNew,rChan,A{rr},rPromo,B{rr})"
    IN[f"D{rr}"] = f"=SUMIFS(rR30,rChan,A{rr},rPromo,B{rr})/C{rr}"
    IN[f"E{rr}"] = f"=SUMIFS(rR90,rChan,A{rr},rPromo,B{rr})/C{rr}"
    IN[f"F{rr}"] = f"=E{rr}/$E$49"
    IN[f"C{rr}"].number_format = NUM
    for col in "DEF":
        IN[f"{col}{rr}"].number_format = PCT
put(IN, "A57", "Promo-acquired vs organic 90-day retention gap by region (DoorDash + UberEats)", f_bold)
hdr(IN, 58, 1, ["Region", "Organic 90d", "Promo 90d", "Promo gap", "Share of FY promo $", "Exit promo % gross"])
for i, g in enumerate(regions):
    rr = 59 + i
    IN[f"A{rr}"] = g
    org = f'(SUMIFS(rR90,rReg,$A{rr},rChan,"DoorDash",rPromo,FALSE)+SUMIFS(rR90,rReg,$A{rr},rChan,"UberEats",rPromo,FALSE))/(SUMIFS(rNew,rReg,$A{rr},rChan,"DoorDash",rPromo,FALSE)+SUMIFS(rNew,rReg,$A{rr},rChan,"UberEats",rPromo,FALSE))'
    pro = org.replace("FALSE", "TRUE")
    IN[f"B{rr}"] = "=" + org; IN[f"C{rr}"] = "=" + pro
    IN[f"D{rr}"] = f"=C{rr}/B{rr}-1"
    IN[f"E{rr}"] = f"=SUMIFS(pPromo,pReg,$A{rr})/SUM(pPromo)"
    IN[f"F{rr}"] = f'=SUMIFS(pPromo,pReg,$A{rr},pWk,">="&EXIT_START)/SUMIFS(pGross,pReg,$A{rr},pWk,">="&EXIT_START)'
    for col in "BCDEF":
        IN[f"{col}{rr}"].number_format = PCT
put(IN, "A65", "Note: the data does not say whether retention means return to Sola (any channel) or to the same "
    "platform. Grubhub has no promo cohort. Retention for 3PD-acquired customers is ~half of dine-in either way.", f_note)

# Scenarios
put(IN, "A68", "E. Incrementality-adjusted economic value (annualized)", f_h2)
put(IN, "A69", "Economic value = CM2 − (1 − incrementality) x gross x dine-in CM%. A cannibalized $ of 3PD replaces "
    "a $ that would have earned ~48% in-store.", f_note)
hdr(IN, 70, 1, ["Metric", "Total", "DoorDash", "UberEats", "Grubhub", "East Coast"])
scol = {"B": "B", "C": "C", "D": "D", "E": "E", "F": "M"}  # Platform_Market columns
srows = [("FY annualized gross", "={pm}{g}", USD),
         ("FY annualized CM2", "={pm}{c}", USD),
         ("FY CM2 %", "={pm}{cp}", PCT),
         ("Exit-period CM2 %", "={pm}{ecp}", PCT),
         ("Break-even incrementality (FY)", "=1-{x}73/DINEIN_CM", PCT),
         ("Break-even incrementality (exit)", "=1-{x}74/DINEIN_CM", PCT)]
for i, (lab, f, fm) in enumerate(srows):
    rr = 71 + i
    put(IN, f"A{rr}", lab)
    for col, pmc in scol.items():
        IN[f"{col}{rr}"] = f.format(pm=f"Platform_Market!{pmc}", g=ROW["FY annualized gross"], c=ROW["FY annualized CM2"],
                                    cp=ROW["CM2 % of gross"], ecp=ROW["Exit period: CM2 % of gross"], x=col)
        IN[f"{col}{rr}"].number_format = fm
        if i < 4:
            IN[f"{col}{rr}"].font = f_link
for j, (lab, nm) in enumerate([("Bear", "INC_BEAR"), ("Base", "INC_BASE"), ("Bull", "INC_BULL")]):
    rr = 78 + j
    put(IN, f"A{rr}", f"Economic value — {lab} (incrementality = input)", bold=(lab == "Base"))
    for col in scol:
        IN[f"{col}{rr}"] = f"={col}72-(1-{nm})*{col}71*DINEIN_CM"; IN[f"{col}{rr}"].number_format = USD
    rr2 = 82 + j
    put(IN, f"A{rr2}", f"Economic value at exit CM2% — {lab}", bold=(lab == "Base"))
    for col in scol:
        IN[f"{col}{rr2}"] = f"={col}71*{col}74-(1-{nm})*{col}71*DINEIN_CM"; IN[f"{col}{rr2}"].number_format = USD
put(IN, "A77", "Incrementality inputs: Bear / Base / Bull →", f_note)
IN["B77"] = "=INC_BEAR"; IN["C77"] = "=INC_BASE"; IN["D77"] = "=INC_BULL"
for col in "BCD":
    IN[f"{col}77"].number_format = PCT; IN[f"{col}77"].font = f_link
for rr in (79, 83):
    for col in "ABCDEF":
        IN[f"{col}{rr}"].fill = fill_key

# StoreGrowth helper sheet
SG = wb.create_sheet("StoreGrowth", index=wb.sheetnames.index("Weeks"))
title(SG, "Store-level Q1→Q4 growth (same-store) — helper for cross-sectional check", "")
hdr(SG, 2, 1, ["location_id", "In-store Q1/wk", "In-store Q4/wk", "3PD Q1/wk", "In-store growth", "3PD growth"])
same = [src["Locations_Master"].cell(row=r, column=1).value for r in range(2, 124)]
cnt = {}
for r in range(2, nC + 1):
    cnt[src["Channel_Revenue_Summary"].cell(row=r, column=1).value] = cnt.get(src["Channel_Revenue_Summary"].cell(row=r, column=1).value, 0) + 1
same = [s for s in same if cnt[s] == 12]
assert len(same) == 112
for i, s in enumerate(same):
    rr = 3 + i
    SG[f"A{rr}"] = s
    SG[f"B{rr}"] = f'=SUMIFS(cInStore,cLoc,A{rr},cQ,"Q1 Jul-Sep")/13'
    SG[f"C{rr}"] = f'=SUMIFS(cInStore,cLoc,A{rr},cQ,"Q4 Apr-Jun")/12'
    SG[f"D{rr}"] = f'=SUMIFS(c3PD,cLoc,A{rr},cQ,"Q1 Jul-Sep")/13'
    SG[f"E{rr}"] = f"=C{rr}/B{rr}-1"
    SG[f"F{rr}"] = f'=(SUMIFS(c3PD,cLoc,A{rr},cQ,"Q4 Apr-Jun")/12)/D{rr}-1'
    for col in "BCD":
        SG[f"{col}{rr}"].number_format = USD
    SG[f"E{rr}"].number_format = PCT; SG[f"F{rr}"].number_format = PCT
put(SG, "H2", "Store list = stores with 12 months in Channel_Revenue_Summary (QA_Checks confirms 112).", f_note)
widths(SG, {c: 14 for c in "ABCDEF"})

# =====================================================================
# Markup
# =====================================================================
MK = wb["Markup"]
title(MK, "Menu markup test — contribution, not just net revenue",
      "32 DoorDash stores, 8 per arm. README: net revenue per store = volume idx x net/order idx. Indices relative to "
      "0% arm. Gain valued per baseline order using each platform's own FY per-order economics.")
widths(MK, {"A": 44, **{CL(i): 14 for i in range(2, 10)}})
hdr(MK, 4, 1, ["Metric", "0% (control)", "5%", "10%", "15%"])
lv = [0, 5, 10, 15]
blocks = {0: (2, 9), 5: (10, 17), 10: (18, 25), 15: (26, 33)}
mrows = ["Stores (n)", "Mean volume index", "Mean net revenue / order index", "Mean net revenue / store index",
         "Std dev net revenue / store index", "Volume vs control", "Net / order vs control", "Net revenue / store vs control",
         "t-stat: net revenue / store vs control", "Block integrity check (rows = this arm)"]
for i, lab in enumerate(mrows):
    put(MK, f"A{5+i}", lab)
for j, m in enumerate(lv):
    col = CL(2 + j); a, b = blocks[m]
    MK[f"{col}5"] = f"=COUNTIF(Data_Markup!$C$2:$C$33,{m})"
    MK[f"{col}6"] = f"=AVERAGEIF(Data_Markup!$C$2:$C$33,{m},Data_Markup!$D$2:$D$33)"
    MK[f"{col}7"] = f"=AVERAGEIF(Data_Markup!$C$2:$C$33,{m},Data_Markup!$E$2:$E$33)"
    MK[f"{col}8"] = f"=AVERAGEIF(Data_Markup!$C$2:$C$33,{m},Data_Markup!$G$2:$G$33)"
    MK[f"{col}9"] = f"=STDEV(Data_Markup!$G${a}:$G${b})"
    MK[f"{col}10"] = f"={col}6/$B$6"
    MK[f"{col}11"] = f"=({col}8/{col}6)/($B$8/$B$6)"
    MK[f"{col}12"] = f"={col}8/$B$8"
    MK[f"{col}13"] = f"=({col}8-$B$8)/SQRT({col}9^2/{col}5+$B$9^2/$B$5)" if m else None
    MK[f"{col}14"] = f"=COUNTIF(Data_Markup!$C${a}:$C${b},{m})=8"
    for rr, fm in zip(range(5, 15), [NUM, DEC3, DEC3, DEC3, DEC3, PCT, PCT, PCT, "0.00", None]):
        if fm:
            MK[f"{col}{rr}"].number_format = fm
for rr in (10, 11, 12):
    for col in "BCDE":
        MK[f"{col}{rr}"] = f"={MK[f'{col}{rr}'].value[1:]}-1"
        MK[f"{col}{rr}"].number_format = PCT
MK["A10"] = "Volume vs control (Δ)"; MK["A11"] = "Net / order vs control (Δ)"; MK["A12"] = "Net revenue / store vs control (Δ)"
put(MK, "A16", "Volume elasticity, log-log slope across 32 stores"); MK["B16"] = "=SLOPE(Data_Markup!H2:H33,Data_Markup!I2:I33)"
put(MK, "A17", "R² of that fit"); MK["B17"] = "=RSQ(Data_Markup!H2:H33,Data_Markup!I2:I33)"
MK["B16"].number_format = "0.00"; MK["B17"].number_format = "0.00"
put(MK, "C16", "≈ −1: each 1% of markup costs ~1% of orders — but net/order rises more, and COGS/labor don't rise with price.", f_note)

put(MK, "A19", "Contribution impact per platform (annualized) — Δ CM2 per baseline order = vol x (net/order x (1+Δnet/ord) − var cost/order) − (net/order − var cost/order)", f_h2)
hdr(MK, 20, 1, ["Metric", "DoorDash", "UberEats", "Grubhub", "Total"])
pmcol = {"B": "C", "C": "D", "D": "E"}
put(MK, "A21", "FY annualized orders")
put(MK, "A22", "Net revenue per order (FY)")
put(MK, "A23", "Variable cost per order (COGS+labor on un-marked-up gross)")
put(MK, "A24", "CM2 per order (FY)")
for col, pc in pmcol.items():
    MK[f"{col}21"] = f"=Platform_Market!{pc}9*ANN"; MK[f"{col}21"].font = f_link
    MK[f"{col}22"] = f"=Platform_Market!{pc}17/Platform_Market!{pc}9"
    MK[f"{col}23"] = f"=Platform_Market!{pc}10/Platform_Market!{pc}9*(COGS_PCT+LABOR_PCT)"
    MK[f"{col}24"] = f"={col}22-{col}23"
    MK[f"{col}21"].number_format = NUM
    for rr in (22, 23, 24):
        MK[f"{col}{rr}"].number_format = USD2
MK["E21"] = "=SUM(B21:D21)"; MK["E21"].number_format = NUM
for j, m in enumerate([5, 10, 15]):
    rr = 26 + j
    tc = CL(3 + j)  # C/D/E columns in top table for 5/10/15
    put(MK, f"A{rr}", f"Δ annual CM2 at {m}% markup (test-implied)")
    for col in pmcol:
        MK[f"{col}{rr}"] = f"={col}$21*((1+${tc}$10)*({col}$22*(1+${tc}$11)-{col}$23)-{col}$24)"
        MK[f"{col}{rr}"].number_format = USD
    MK[f"E{rr}"] = f"=SUM(B{rr}:D{rr})"; MK[f"E{rr}"].number_format = USD
put(MK, "A30", "Selected markup (Assumptions!MARKUP_SEL)", bold=True); MK["B30"] = "=MARKUP_SEL"; MK["B30"].font = f_link
put(MK, "A31", "Test-implied Δ CM2, selected level")
MK["E31"] = "=INDEX(E26:E28,MATCH(MARKUP_SEL,{5;10;15},0))"
put(MK, "A32", "Banked Δ CM2 after haircut", bold=True); MK["E32"] = "=E31*(1-MARKUP_HAIRCUT)"
put(MK, "A33", "Volume change at selected level"); MK["E33"] = "=INDEX(C10:E10,MATCH(MARKUP_SEL,{5,10,15},0))"
for rr in (31, 32):
    MK[f"E{rr}"].number_format = USD
MK["E33"].number_format = PCT
MK["E32"].fill = fill_key; MK["E32"].font = f_bold
put(MK, "A34", "Cross-check: mechanical Δ CM2 at selected level (price x (1 − %-of-gross fees), test volume)")
MK["E34"] = ("=SUMPRODUCT(B21:D21,(1+E33)*(B22:D22+(MARKUP_SEL/100)*(Platform_Market!C10:E10/Platform_Market!C9:E9)"
             "*(1-(Platform_Market!C11:E11+Platform_Market!C12:E12+Platform_Market!C15:E15+Platform_Market!C16:E16)/Platform_Market!C10:E10)-B23:D23)-B24:D24)")
MK["E34"].number_format = USD
put(MK, "A36", "Mechanical check uses no test net/order data — only the price change and observed volume; banked value should sit at or below it.", f_note)
put(MK, "A35", "Caveats: test is DoorDash-only; applying the same response to UberEats/Grubhub is an ANALYST assumption. "
    "Lost orders are treated as fully lost (none shift to dine-in) — conservative. Markup gain is available with or "
    "without the DoorDash contract, so it must not be credited to 'accept'.", f_note)

# =====================================================================
# DoorDash
# =====================================================================
DD = wb["DoorDash"]
title(DD, "DoorDash contract — accept vs counter vs walk",
      "Year-1 base = model-derived annualized DoorDash gross (vs ~$39.5M in Contract_Terms). Year 2 grows at DD_GROWTH. "
      "Costs are CM2 impact vs status quo (positive = cost to Sola).")
widths(DD, {"A": 50, "B": 16, "C": 16, "D": 16, "E": 16, "F": 16, "G": 60})
put(DD, "A4", "Current DoorDash economics (FY, annualized)", f_h2)
dd_in = [
    ("Gross (51-wk actual)", "=Platform_Market!C10", USD),
    ("Year-1 gross (annualized)", "=B5*ANN", USD),
    ("Year-2 gross", "=B6*(1+DD_GROWTH)", USD),
    ("Observed commission rate", "=Platform_Market!C11/Platform_Market!C10", PCT2),
    ("Promo (annualized)", "=Platform_Market!C13*ANN", USD),
    ("Ads (annualized)", "=Platform_Market!C14*ANN", USD),
    ("Current ad share of DoorDash marketing", "=B10/(B9+B10)", PCT),
    ("CM2 % (FY)", "=Platform_Market!C26", PCT),
    ("Contract_Terms stated run-rate (text '~$39.5M')", 39500000, USD),
    ("Model Year-1 gross vs stated run-rate", "=B6/B13-1", PCT),
]
for i, (lab, f, fm) in enumerate(dd_in):
    rr = 5 + i
    put(DD, f"A{rr}", lab)
    c = DD[f"B{rr}"]; c.value = f; c.number_format = fm
    c.font = f_in if not isinstance(f, str) else (f_link if "Platform_Market" in f else f_base)
    if not isinstance(f, str):
        c.fill = fill_in
put(DD, "G13", "Hardcoded from Contract_Terms text (cross-check only; not used in calculations).", f_note)
# NOTE: rows below are referenced by Promo_Escalation (DoorDash!C12) -> keep layout stable
put(DD, "A11", "Current ad share of DoorDash marketing")  # (row 11 already)
hdr(DD, 15, 1, ["Scenario (annual CM2 cost vs status quo)", "Status quo", "Accept — spend-neutral mix",
                "Accept — ads top-up", "Counter", "Walk", "Notes"])
# Put headline commission delta row 12 col C for reference
put(DD, "A16", "Headline commission")
DD["B16"] = "=DD_COMM_CUR"; DD["C16"] = "=DD_COMM_PROP"; DD["D16"] = "=DD_COMM_PROP"; DD["E16"] = "=CTR_COMM"; DD["F16"] = "n/a"
put(DD, "A17", "Required ads share of DoorDash marketing")
DD["B17"] = "=B11"; DD["C17"] = "=DD_ADS_MIN"; DD["D17"] = "=DD_ADS_MIN"; DD["E17"] = "=CTR_ADS"
put(DD, "A18", "Term (months)")
DD["B18"] = 12; DD["C18"] = 24; DD["D18"] = 24; DD["E18"] = 12; DD["F18"] = "—"
for col in "BCDE":
    DD[f"{col}18"].font = f_in
put(DD, "A19", "Y1 incremental commission"); put(DD, "A20", "Y2 incremental commission")
put(DD, "A21", "Y1 mandate cost (ads top-up, holding FY promo)"); put(DD, "A22", "Y2 mandate cost")
put(DD, "A23", "Y1 lost CM2 (walk)"); put(DD, "A24", "Y1 recaptured CM2 (walk)")
put(DD, "A25", "12-month cost", bold=True); put(DD, "A26", "24-month cost", bold=True)
put(DD, "A27", "DoorDash CM2 % after (Y1)"); put(DD, "A28", "DoorDash all-in take after (FY mix)")
put(DD, "A29", "Memo: UberEats FY CM2 % (worst platform today)")
for col in "BCDE":
    DD[f"{col}19"] = f"=({col}16-DD_COMM_CUR)*$B$6"
    DD[f"{col}20"] = f"=({col}16-DD_COMM_CUR)*$B$7"
DD["B21"] = 0; DD["C21"] = 0; DD["E21"] = 0
DD["D21"] = "=MAX(0,$B$9*D17/(1-D17)-$B$10)"
DD["B22"] = 0; DD["C22"] = 0; DD["E22"] = 0
DD["D22"] = "=D21*(1+DD_GROWTH)"
for col in "BCDE":
    DD[f"{col}23"] = 0; DD[f"{col}24"] = 0
DD["F19"] = 0; DD["F20"] = 0; DD["F21"] = 0; DD["F22"] = 0
DD["F23"] = "=$B$6*$B$12"
DD["F24"] = "=-$B$6*(WALK_R3*Platform_Market!D26*Platform_Market!D10/(Platform_Market!D10+Platform_Market!E10)+WALK_R3*Platform_Market!E26*Platform_Market!E10/(Platform_Market!D10+Platform_Market!E10)+WALK_ROWN*DINEIN_CM)"
for col in "BCDEF":
    DD[f"{col}25"] = f"={col}19+{col}21+{col}23+{col}24"
DD["B26"] = "=B25*2"
DD["C26"] = "=C19+C20+C21+C22"; DD["D26"] = "=D19+D20+D21+D22"
DD["E26"] = "=E19+E20+E21+E22"
DD["F26"] = "=F25+(F23+F24)*(1+DD_GROWTH)"
for col in "BCDE":
    DD[f"{col}27"] = f"=$B$12-({col}19+{col}21)/$B$6"
    DD[f"{col}28"] = f"=Platform_Market!$C$18+({col}19+{col}21)/$B$6"
DD["B29"] = "=Platform_Market!D26"; DD["B29"].font = f_link
DD["B25"] = 0; DD["B26"] = 0
notes = {19: "Δ commission x Y1 gross", 20: "Δ commission x Y2 gross (8% growth)",
         21: "Ads needed so ads ≥ 60% while holding FY promo; spend-neutral case reallocates instead (cost 0, efficiency unknown)",
         23: "Walk loses DoorDash FY CM2", 24: "Recapture: WALK_R3 to UE/GH at their CM2%, WALK_ROWN to own channels at 48%",
         26: "Counter assumes same terms renewed for Y2 (comparable horizon)"}
for k, v in notes.items():
    put(DD, f"G{k}", v, f_note)
for rr in range(19, 27):
    for col in "BCDEF":
        DD[f"{col}{rr}"].number_format = USD
for col in "BCDE":
    DD[f"{col}16"].number_format = PCT; DD[f"{col}17"].number_format = PCT
    DD[f"{col}27"].number_format = PCT; DD[f"{col}28"].number_format = PCT
DD["B29"].number_format = PCT
for col in "BCDEF":
    DD[f"{col}25"].fill = fill_key; DD[f"{col}26"].fill = fill_key
    DD[f"{col}25"].font = f_bold; DD[f"{col}26"].font = f_bold
# Reference cell used by Promo_Escalation: C12 = annual cost of +4pt ask
put(DD, "C11", "← memo", f_note)
DD["C12"] = "=(DD_COMM_PROP-DD_COMM_CUR)*B6"; DD["C12"].number_format = USD
put(DD, "D12", "Y1 cost of +4 pt ask (memo)", f_note)

put(DD, "A31", "Decision thresholds", f_h2)
put(DD, "A32", "Walk cost, Y1 (net of recapture)"); DD["B32"] = "=F25"
put(DD, "A33", "Break-even commission vs walking — spend-neutral mix"); DD["B33"] = "=DD_COMM_CUR+B32/B6"
put(DD, "A34", "Break-even commission vs walking — with ads top-up"); DD["B34"] = "=DD_COMM_CUR+(B32-D21)/B6"
put(DD, "A35", "Savings: counter vs accept (spend-neutral), 24 months"); DD["B35"] = "=C26-E26"
put(DD, "A36", "Savings: counter vs accept (top-up), 24 months"); DD["B36"] = "=D26-E26"
put(DD, "A37", "Markup on DoorDash needed to offset +4 pts at flat volume")
DD["B37"] = "=(DD_COMM_PROP-DD_COMM_CUR)/(1-DD_COMM_PROP-Platform_Market!C12/Platform_Market!C10-Platform_Market!C15/Platform_Market!C10-Platform_Market!C16/Platform_Market!C10)"
put(DD, "G37", "Percentage-of-gross deductions scale with price; promo/ads held in $. Markup is available regardless of the contract.", f_note)
put(DD, "A38", "Walk cost if own-channel recapture = 1 − Base incrementality (consistency test)")
DD["B38"] = "=$B$6*$B$12-$B$6*(WALK_R3*(Platform_Market!D25+Platform_Market!E25)/(Platform_Market!D10+Platform_Market!E10)+(1-INC_BASE)*DINEIN_CM)"
put(DD, "G38", "If Base incrementality held, walking would roughly break even or better — but that rests on one weak inference and ignores peak-seat capacity.", f_note)
for rr in (32, 35, 36, 38):
    DD[f"B{rr}"].number_format = USD
for rr in (33, 34, 37):
    DD[f"B{rr}"].number_format = PCT2
    DD[f"B{rr}"].fill = fill_key

put(DD, "A39", "Ads top-up still needed for 60% mandate AFTER promo rollback (Y1)")
DD["B39"] = "=MAX(0,Promo_Escalation!B6*$B$6*DD_ADS_MIN/(1-DD_ADS_MIN)-$B$10)"
DD["B39"].number_format = USD
put(DD, "G39", "With promo back at baseline (DoorDash ~1.3% of gross), the 60% ads floor needs only a small shift — the mandate is cheap if Sola fixes promo first.", f_note)
put(DD, "A41", "Sensitivity: Y1 walk cost by recapture assumptions (rows = to UE/GH; cols = to own channels)", f_h2)
own = [0.0, 0.10, 0.15, 0.20, 0.30, 0.45]
r3 = [0.0, 0.15, 0.30, 0.45, 0.60]
DD["A42"] = "UE/GH ↓  /  own →"
for j, o in enumerate(own):
    c = DD.cell(row=42, column=2 + j, value=o); c.number_format = PCT; c.font = f_in
for i, x in enumerate(r3):
    rr = 43 + i
    c = DD.cell(row=rr, column=1, value=x); c.number_format = PCT; c.font = f_in
    for j in range(len(own)):
        col = CL(2 + j)
        DD[f"{col}{rr}"] = (f"=$B$6*$B$12-$B$6*($A{rr}*(Platform_Market!$D$25+Platform_Market!$E$25)/"
                            f"(Platform_Market!$D$10+Platform_Market!$E$10)+{col}$42*DINEIN_CM)")
        DD[f"{col}{rr}"].number_format = USD
put(DD, "A49", "Compare each cell with the 12-month accept cost (C25 / D25): a walk cost below it means walking is cheaper.", f_note)
put(DD, "A50", "Note: grid uses blended UE+GH CM2 (CM2 $/gross $); scenario row 24 weights the two platforms' CM2% by gross — identical by construction.", f_note)

# =====================================================================
# Sensitivity
# =====================================================================
SE = wb["Sensitivity"]
title(SE, "Sensitivity analysis", "Live formula grids (no Data Tables) — every cell recomputes from model inputs.")
widths(SE, {"A": 34, **{CL(i): 13 for i in range(2, 10)}})
put(SE, "A4", "1. FY channel CM2 % by COGS % (rows) and variable labor % (cols)", f_h2)
cogs = [0.26, 0.28, 0.30, 0.32, 0.34]; lab_ = [0.12, 0.14, 0.16, 0.18, 0.20]
for j, v in enumerate(lab_):
    c = SE.cell(row=5, column=2 + j, value=v); c.number_format = PCT; c.font = f_in
for i, v in enumerate(cogs):
    rr = 6 + i
    c = SE.cell(row=rr, column=1, value=v); c.number_format = PCT; c.font = f_in
    for j in range(5):
        col = CL(2 + j)
        SE[f"{col}{rr}"] = f"=PL_Waterfall!$B$13/PL_Waterfall!$B$6-$A{rr}-{col}$5"; SE[f"{col}{rr}"].number_format = PCT
put(SE, "A12", "2. FY economic value of the channel ($) by incrementality (rows) and dine-in CM% (cols)", f_h2)
inc = [0.3, 0.45, 0.55, 0.62, 0.7, 0.8, 0.9]; dcm = [0.40, 0.44, 0.48, 0.52, 0.56]
for j, v in enumerate(dcm):
    c = SE.cell(row=13, column=2 + j, value=v); c.number_format = PCT; c.font = f_in
for i, v in enumerate(inc):
    rr = 14 + i
    c = SE.cell(row=rr, column=1, value=v); c.number_format = PCT; c.font = f_in
    for j in range(5):
        col = CL(2 + j)
        SE[f"{col}{rr}"] = f"=PL_Waterfall!$E$18-(1-$A{rr})*PL_Waterfall!$E$6*{col}$13"; SE[f"{col}{rr}"].number_format = USD
put(SE, "A22", "3. DoorDash 24-month cost of a commission increase ($) by headline rate (rows) and Y2 growth (cols) — spend-neutral mix", f_h2)
cr = [0.28, 0.285, 0.29, 0.30, 0.31]; gr = [0.0, 0.04, 0.08, 0.12, 0.16]
for j, v in enumerate(gr):
    c = SE.cell(row=23, column=2 + j, value=v); c.number_format = PCT; c.font = f_in
for i, v in enumerate(cr):
    rr = 24 + i
    c = SE.cell(row=rr, column=1, value=v); c.number_format = PCT; c.font = f_in
    for j in range(5):
        col = CL(2 + j)
        SE[f"{col}{rr}"] = f"=($A{rr}-DD_COMM_CUR)*DoorDash!$B$6*(2+{col}$23)"; SE[f"{col}{rr}"].number_format = USD
put(SE, "A30", "4. Promo rollback net CM2 ($/yr) by share of escalation removed (rows) and volume loss (cols)", f_h2)
sh = [0.25, 0.5, 0.75, 1.0]; vl = [0.0, 0.03, 0.05, 0.10, 0.15]
for j, v in enumerate(vl):
    c = SE.cell(row=31, column=2 + j, value=v); c.number_format = PCT; c.font = f_in
for i, v in enumerate(sh):
    rr = 32 + i
    c = SE.cell(row=rr, column=1, value=v); c.number_format = PCT; c.font = f_in
    for j in range(5):
        col = CL(2 + j)
        SE[f"{col}{rr}"] = (f"=Promo_Escalation!$B$22*$A{rr}*(1-{col}$31)-Promo_Escalation!$B$20*{col}$31*PL_Waterfall!$G$18")
        SE[f"{col}{rr}"].number_format = USD
put(SE, "A37", "Volume loss here is gross volume lost because promo is withdrawn; it is applied at exit CM2%.", f_note)

# =====================================================================
# Exec_Summary
# =====================================================================
ES = wb["Exec_Summary"]
title(ES, "Executive summary — every number links to the model",
      "Objective optimized: MARGIN (CM2 dollars / EBITDA quality ahead of a raise). Recommendation: HOLD — reprice, don't retreat.")
widths(ES, {"A": 64, "B": 18, "C": 18, "D": 70})
hdr(ES, 4, 1, ["Metric", "FY (annualized)", "Exit run-rate", "Source"])
es = [
    ("3PD gross sales", "=PL_Waterfall!E6", "=PL_Waterfall!F6", USD, "PL_Waterfall row 6"),
    ("Net revenue (bank deposit)", "=PL_Waterfall!E13", "=PL_Waterfall!F13", USD, "row 13"),
    ("All-in take rate", "=PL_Waterfall!C14", "=PL_Waterfall!G14", PCT, "row 14"),
    ("CM1", "=PL_Waterfall!E16", "=PL_Waterfall!F16", USD, "row 16"),
    ("CM2 (contribution)", "=PL_Waterfall!E18", "=PL_Waterfall!F18", USD, "row 18"),
    ("CM2 % of gross", "=PL_Waterfall!C18", "=PL_Waterfall!G18", PCT, "row 18"),
    ("Dine-in benchmark", "=DINEIN_CM", "=DINEIN_CM", PCT, "Assumption"),
    ("Break-even incrementality", "=PL_Waterfall!C23", "=PL_Waterfall!G23", PCT, "row 23"),
    ("CM2 % — DoorDash", "=Platform_Market!C26", f"=Platform_Market!C{ROW['Exit period: CM2 % of gross']}", PCT, "Platform_Market"),
    ("CM2 % — UberEats", "=Platform_Market!D26", f"=Platform_Market!D{ROW['Exit period: CM2 % of gross']}", PCT, ""),
    ("CM2 % — Grubhub", "=Platform_Market!E26", f"=Platform_Market!E{ROW['Exit period: CM2 % of gross']}", PCT, ""),
    ("CM2 % — Established markets", "=Platform_Market!L26", f"=Platform_Market!L{ROW['Exit period: CM2 % of gross']}", PCT, ""),
    ("CM2 % — East Coast (new)", "=Platform_Market!M26", f"=Platform_Market!M{ROW['Exit period: CM2 % of gross']}", PCT, ""),
    ("Promo % of gross (pre-Feb → exit)", "=Promo_Escalation!B5", "=Promo_Escalation!C5", PCT2, "Promo_Escalation"),
    ("Annual cost of promo escalation at exit run-rate", None, "=Promo_Escalation!B22", USD, "Promo_Escalation B22"),
    ("Established gross / store-wk (late-Jan → exit)", "=Promo_Escalation!B9", "=Promo_Escalation!C9", USD, ""),
    ("Implied incrementality of marginal 3PD growth", "=Incrementality!B30", None, PCT, "Incrementality B30"),
    ("3PD orders in peak dayparts (seats 70–95% full)", "=Incrementality!H43", None, PCT, "Incrementality H43"),
    ("90-day retention: dine-in vs DoorDash (organic)", "=Incrementality!E49", "=Incrementality!E51", PCT, "B = dine-in, C = DoorDash"),
]
for i, (lab, f1, f2, fm, s) in enumerate(es):
    rr = 5 + i
    put(ES, f"A{rr}", lab)
    for col, f in (("B", f1), ("C", f2)):
        if f:
            ES[f"{col}{rr}"] = f; ES[f"{col}{rr}"].number_format = fm; ES[f"{col}{rr}"].font = f_link
    put(ES, f"D{rr}", s, f_note)
rr = 5 + len(es) + 1
put(ES, f"A{rr}", "DoorDash decision (annual / 24-month CM2 cost vs status quo)", f_h2); rr += 1
ESD = rr
dd_es = [("Accept 31% — spend-neutral ad mix", "=DoorDash!C25", "=DoorDash!C26"),
         ("Accept 31% — ads top-up to 60%", "=DoorDash!D25", "=DoorDash!D26"),
         ("Counter 28.5% / ≤40% ads / 12-mo (RECOMMENDED)", "=DoorDash!E25", "=DoorDash!E26"),
         ("Walk (base recapture)", "=DoorDash!F25", "=DoorDash!F26"),
         ("Break-even commission vs walk (spend-neutral / top-up)", "=DoorDash!B33", "=DoorDash!B34")]
hdr(ES, rr, 1, ["Option", "12-month", "24-month", ""]); rr += 1
for lab, f1, f2 in dd_es:
    put(ES, f"A{rr}", lab, bold="RECOMMENDED" in lab)
    ES[f"B{rr}"] = f1; ES[f"C{rr}"] = f2
    fm = PCT2 if "Break-even" in lab else USD
    ES[f"B{rr}"].number_format = fm; ES[f"C{rr}"].number_format = fm
    ES[f"B{rr}"].font = f_link; ES[f"C{rr}"].font = f_link
    rr += 1
rr += 1
put(ES, f"A{rr}", "Recommended plan — annual CM2 impact vs exit run-rate", f_h2); rr += 1
hdr(ES, rr, 1, ["Action", "Annual CM2 ($)", "Confidence", "Basis"]); rr += 1
PLAN0 = rr
plan = [("1. Roll promo back to pre-Feb-2026 level (~1.5% of gross)", "=Promo_Escalation!B25", "Medium-high",
         "Promo up ~5x with flat/declining gross per store-week; 3% volume-loss haircut"),
        ("2. 5% 3PD menu markup (banked at 50% of test result)", "=Markup!E32", "Medium-low",
         "DoorDash-only test, n=8/arm, noisy; extended to UE/GH by assumption"),
        ("3. Counter DoorDash at 28.5% (cost vs today's terms)", "=-DoorDash!E25", "Medium",
         "Negotiation outcome uncertain; ceiling = accept 31% spend-neutral"),
        ]
for lab, f, cf, b in plan:
    put(ES, f"A{rr}", lab); ES[f"B{rr}"] = f; ES[f"B{rr}"].number_format = USD; ES[f"B{rr}"].font = f_link
    put(ES, f"C{rr}", cf); put(ES, f"D{rr}", b, f_note); rr += 1
put(ES, f"A{rr}", "Net annual CM2 improvement vs exit run-rate", bold=True)
ES[f"B{rr}"] = f"=SUM(B{PLAN0}:B{rr-1})"; ES[f"B{rr}"].number_format = USD; ES[f"B{rr}"].font = f_bold; ES[f"B{rr}"].fill = fill_key
PLANTOT = rr; rr += 1
put(ES, f"A{rr}", "Resulting CM2 run-rate"); ES[f"B{rr}"] = f"=PL_Waterfall!F18+B{PLANTOT}"; ES[f"B{rr}"].number_format = USD
rr += 1
put(ES, f"A{rr}", "Resulting CM2 % of exit gross (approx.; ignores markup gross uplift)")
ES[f"B{rr}"] = f"=B{rr-1}/PL_Waterfall!F6"; ES[f"B{rr}"].number_format = PCT
rr += 2
put(ES, f"A{rr}", "If the board optimizes for GROWTH instead — 'lean in selectively' variant", f_h2); rr += 1
GV0 = rr
gv = [("Promo rollback only where promo cohorts retain badly (TX, SE, Mountain West)",
       "=Promo_Escalation!B25*(Incrementality!E59+Incrementality!E60+Incrementality!E62)",
       "Keeps promo in SoCal/PNW/Northeast where promo-acquired 90d retention is within ~10–13% of organic"),
      ("5% markup (same as margin plan)", "=Markup!E32", "Volume −1.25% in test"),
      ("Accept DoorDash 31% / 24-mo, spend-neutral mix (placement & lock-in)", "=-DoorDash!C25", "Secures placement for East Coast growth")]
for lab, f, b in gv:
    put(ES, f"A{rr}", lab); ES[f"B{rr}"] = f; ES[f"B{rr}"].number_format = USD; ES[f"B{rr}"].font = f_link
    put(ES, f"D{rr}", b, f_note); rr += 1
put(ES, f"A{rr}", "Growth variant: net annual CM2 vs exit run-rate", bold=True)
ES[f"B{rr}"] = f"=SUM(B{GV0}:B{rr-1})"; ES[f"B{rr}"].number_format = USD; ES[f"B{rr}"].font = f_bold
rr += 1
put(ES, f"A{rr}", "CM2 given up vs margin plan (price of the growth objective)", bold=True)
ES[f"B{rr}"] = f"=B{rr-1}-B{PLANTOT}"; ES[f"B{rr}"].number_format = USD; ES[f"B{rr}"].fill = fill_key
rr += 2
put(ES, f"A{rr}", "Interactions between actions are not modeled (e.g., markup on a lower-promo base). Totals are indicative.", f_note)

# =====================================================================
# QA_Checks
# =====================================================================
QA = wb["QA_Checks"]
title(QA, "QA checks — all should read TRUE", "Recomputed by LibreOffice headless; see QA_Summary.md for results.")
widths(QA, {"A": 70, "B": 18, "C": 18, "D": 10})
hdr(QA, 4, 1, ["Check", "Value", "Expected", "Pass"])
_regs = ["Texas", "Southeast", "Southern California", "Mountain West", "Pacific Northwest", "Northeast"]
DP_FLAG = "=MAX(" + ",".join(f'ABS(SUMIF(Data_Daypart!$A$2:$A$25,"{g}",Data_Daypart!${c}$2:${c}$25)-1)' for g in _regs for c in "CDE") + ")"
qa = [
    ("Payout rows", f"=COUNTA(Data_Payouts!A2:A{nP})", 17934, "eq"),
    ("Unique weeks", "=WEEKS_IN_DATA", 51, "eq"),
    ("Stores in Locations_Master", "=COUNTA(Data_Locations!A2:A123)", 122, "eq"),
    ("Max |row recon diff| gross − deductions − net (≤ $0.05)", "=MAX(MAX(pDiff),-MIN(pDiff))", 0.05, "le"),
    ("Rows with unmatched location (market_status lookup errors)", f"=SUMPRODUCT(--ISERROR(Data_Payouts!Q2:Q{nP}))", 0, "eq"),
    ("Weeks table total gross = Data_Payouts gross", "=ROUND(Weeks!G57-SUM(pGross),2)", 0, "eq"),
    ("Periods total gross = Data_Payouts gross", "=ROUND(Periods!D19-SUM(pGross),2)", 0, "eq"),
    ("Periods weeks sum = 51", "=Periods!B19", 51, "eq"),
    ("Latest period has 4 weeks", "=Periods!B18", 4, "eq"),
    ("P&L net (computed) − source net deposit, $ (cumulative source penny rounding; |x| ≤ 10)", "=ABS(PL_Waterfall!B26)", 10, "le"),
    ("Exit run-rate net − source x13, $ (|x| ≤ 25)", "=ABS(PL_Waterfall!F26)", 25, "le"),
    ("P&L CM2 vs independent recompute from source net, $ (|x| ≤ 10)", "=ABS(PL_Waterfall!B18-(SUM(pNet)-SUM(pGross)*(COGS_PCT+LABOR_PCT)))", 10, "le"),
    ("Platforms sum to total CM2", f"=Platform_Market!B{PMCHK+1}", 0, "eq"),
    ("Regions sum to total CM2", f"=Platform_Market!B{PMCHK+3}", 0, "eq"),
    ("Status sum to total CM2", f"=Platform_Market!B{PMCHK+4}", 0, "eq"),
    ("Platform x status sum to total CM2", f"=Platform_Market!B{PMCHK+5}", 0, "eq"),
    ("Channel_Revenue 3PD gross = payouts gross (same source)", "=ROUND(SUM(c3PD)-SUM(pGross),2)", 0, "eq"),
    ("Channel_Revenue components sum to total_revenue (max abs diff ≤ $0.05)",
     f"=SUMPRODUCT(ABS(Data_Channel!G2:G{nC}+Data_Channel!H2:H{nC}+Data_Channel!I2:I{nC}+Data_Channel!J2:J{nC}+Data_Channel!K2:K{nC}-Data_Channel!L2:L{nC})>0.05)", 0, "eq"),
    ("Same-store count (12 months of data)", "=SUMPRODUCT(--Data_Channel!M2:M1409)/12", 112, "eq"),
    ("DoorDash observed commission = 27.0% (Contract_Terms)", "=ROUND(DoorDash!B8,3)", 0.27, "eq"),
    ("UberEats observed commission = 30.5%", "=ROUND(Platform_Market!D11/Platform_Market!D10,3)", 0.305, "eq"),
    ("Grubhub observed commission = 22.5%", "=ROUND(Platform_Market!E11/Platform_Market!E10,3)", 0.225, "eq"),
    ("DoorDash marketing commission = 1.5%", "=ROUND(Platform_Market!C12/Platform_Market!C10,3)", 0.015, "eq"),
    ("Marketing spend (promo+ads) 51-wk ≈ $2.9M (Contract_Terms)", "=ROUND(PL_Waterfall!B28/100000,0)/10", 2.9, "eq"),
    ("Ad share of marketing ≈ 27% (Contract_Terms)", "=ROUND(PL_Waterfall!B29,2)", 0.27, "eq"),
    ("DoorDash ad share ≈ 30% (Contract_Terms 'current')", "=ROUND(DoorDash!B11,1)", 0.3, "eq"),
    ("Markup arms are contiguous blocks of 8", "=AND(Markup!B14,Markup!C14,Markup!D14,Markup!E14)", True, "eq"),
    ("Cohort rows: retention monotone (30d ≥ 60d ≥ 90d) violations",
     f"=SUMPRODUCT((Data_Cohort!J2:J{nR}>Data_Cohort!I2:I{nR})+(Data_Cohort!K2:K{nR}>Data_Cohort!J2:J{nR}))", 0, "eq"),
    ("DATA FLAG (informational): daypart shares — max abs deviation of sum from 1 (normalized in model)", DP_FLAG, 0.10, "le"),
]
for i, (lab, f, exp, op) in enumerate(qa):
    rr = 5 + i
    put(QA, f"A{rr}", lab); QA[f"B{rr}"] = f; QA[f"C{rr}"] = exp
    QA[f"D{rr}"] = f"=B{rr}=C{rr}" if op == "eq" else f"=B{rr}<=C{rr}"
    QA[f"B{rr}"].number_format = "General"
rr = 5 + len(qa)
put(QA, f"A{rr}", "ALL CHECKS PASS", bold=True); QA[f"D{rr}"] = f"=AND(D5:D{rr-1})"; QA[f"D{rr}"].font = f_bold

# =====================================================================
# README
# =====================================================================
RM = wb["README"]
title(RM, "Sola Kitchen — Marketplace (3PD) P&L Model", "Prepared for the board review. Fictional company; synthetic data.")
RM.column_dimensions["A"].width = 30; RM.column_dimensions["B"].width = 120
txt = [
    ("Purpose", "Treat third-party marketplace delivery as a standalone P&L down to CM2, judge it against dine-in, "
                "and size the DoorDash contract decision."),
    ("Objective chosen", "MARGIN — CM2 dollars / EBITDA quality ahead of a possible raise. Growth variant discussed in Exec_Summary/deck."),
    ("Build", "Every output is a formula on the raw data tabs (Data_*), which are copied value-for-value from "
              "Sola_Kitchen_Data.xlsx. Calc columns appended to raw tabs are labelled 'CALC' in a lighter header."),
    ("Flow", "Data_Payouts → Weeks → Periods → PL_Waterfall / Platform_Market → Promo_Escalation, Incrementality, "
             "Markup, DoorDash → Sensitivity → Exec_Summary. QA_Checks validates the chain."),
    ("Period conventions", "Payouts are weekly (51 weeks, Jul 7 2025 – Jun 22 2026). FY = 51-wk actual x 52/51 "
                           "(13 periods x 28 days). Exit run-rate = latest 28-day period (Jun 1–28 2026) x 13. "
                           "Channel_Revenue_Summary months contain 4 or 5 payout weeks (Sep, Dec, Mar = 5) so all "
                           "monthly comparisons are per store-week."),
    ("Sign convention", "Source deductions are positive amounts; P&L shows them negative. Net = gross − commission − "
                        "marketing commission − promo − ads − service fee − refunds (README definition; verified row by row)."),
    ("Exclusions", "tax_collected_remitted is marketplace-facilitator tax: excluded from revenue and deductions (memo only)."),
    ("Not CM2", "Bank deposit is NOT contribution: COGS (30%) and variable labor (16%) are deducted after it."),
    ("Color code", "Blue on yellow = input; black = formula; green = cross-sheet link."),
    ("Classification", "OBSERVED = in the data; ASSIGNMENT = given in the brief/README; ANALYST = my judgment "
                       "(each tagged on Assumptions)."),
    ("Data audit — clean", "No nulls, no duplicate keys (platform x store x week), no negative values; every row "
                           "reconciles to net within $0.03; 3 platforms x 122 stores; region/name labels consistent "
                           "across tabs; Channel_Revenue 3PD gross = payouts gross exactly (same source, not an independent check)."),
    ("Data audit — issues", "(1) 51 weeks, not 52. (2) Locations 'state' is a region code, not a real state (Salt Lake "
                            "City/Tempe/Las Vegas = CO; Portland/Bend = WA; Boston/Hoboken/Stamford = NY). (3) "
                            "months_open does not match a single as-of date. (4) 10 East Coast stores open mid-window "
                            "→ use per-store-week metrics. (5) Contract_Terms values are text (~$39.5M, '27.0%'); "
                            "re-entered as numeric inputs and cross-checked to payouts. (6) Cohort retention "
                            "definition (return to brand vs to platform) not stated. (7) Markup test is DoorDash-only, n=8/arm."),
    ("Key finding to verify", "Restaurant-funded promo is flat at ~1.5% of gross through Feb 2 2026, then rises every "
                              "week to ~7.4% by late June across ALL platforms and regions. Confirm with Marketing whether "
                              "this is a real program, a platform auto-promo, or a data/classification change."),
]
for i, (k, v) in enumerate(txt):
    rr = 4 + i
    put(RM, f"A{rr}", k, bold=True); put(RM, f"B{rr}", v, wrap=True)
    RM.row_dimensions[rr].height = 15 * max(1, len(v) // 110 + 1)
put(RM, f"A{5+len(txt)}", "Sheets", f_h2)
for i, s in enumerate(wb.sheetnames):
    put(RM, f"A{6+len(txt)+i}", s)

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != F:
                c.font = Font(name=F, size=10, bold=c.font.bold if c.font else False,
                              color=c.font.color if c.font else None, italic=c.font.italic if c.font else False)

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
